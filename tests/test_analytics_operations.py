import csv
import io
import json
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from dongjiang_agent.operations import AnalyticsService, task_intervals


class AnalyticsOperationsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.case_root = self.root / "cases"
        self.case_root.mkdir()
        self.policy_path = self.root / "sla.json"
        self.policy = {
            "version": "analytics-test-v1",
            "tasks": {
                "credit_approval": {"label": "信用审批", "target_hours": 24},
                "contract_upload": {"label": "上传合同", "target_hours": 72},
                "sales_revision": {"label": "修改合同", "target_hours": 48},
            },
        }
        self.policy_path.write_text(
            json.dumps(self.policy, ensure_ascii=False), encoding="utf-8"
        )

    def tearDown(self):
        self.temp.cleanup()

    def save_case(self, payload):
        (self.case_root / f"{payload['case_id']}.json").write_text(
            json.dumps(payload, ensure_ascii=False), encoding="utf-8"
        )

    def test_trace_intervals_reconstruct_completed_and_current_tasks(self):
        case = {
            "case_id": "DJ-ANALYTICS1",
            "status": "blocked",
            "created_at": "2026-07-01T08:00:00+00:00",
            "customer": {"customer_name": "分析客户", "business_type": "TKP"},
            "owner": {"user_id": "sales-1", "display_name": "销售甲"},
            "trace": [
                {"ts": "2026-07-01T08:00:00+00:00", "stage": "credit.approval_requested"},
                {"ts": "2026-07-01T18:00:00+00:00", "stage": "credit.effective"},
                {"ts": "2026-07-01T18:00:00+00:00", "stage": "workflow.interrupt"},
                {"ts": "2026-07-02T06:00:00+00:00", "stage": "workflow.resumed"},
                {
                    "ts": "2026-07-02T08:00:00+00:00",
                    "stage": "decision.routed",
                    "data": {"decision": "block"},
                },
            ],
        }
        rows = task_intervals(
            case,
            now=datetime(2026, 7, 4, 20, tzinfo=timezone.utc),
            policy=self.policy,
        )
        self.assertEqual([row["task"] for row in rows], ["credit_approval", "contract_upload", "sales_revision"])
        self.assertEqual(rows[0]["duration_hours"], 10.0)
        self.assertTrue(rows[0]["met_sla"])
        self.assertEqual(rows[1]["duration_hours"], 12.0)
        self.assertFalse(rows[2]["completed"])
        self.assertEqual(rows[2]["duration_hours"], 60.0)
        self.assertFalse(rows[2]["met_sla"])
        self.assertEqual(rows[2]["responsible"]["label"], "销售甲")

    def test_report_filters_by_case_start_and_aggregates_metrics(self):
        self.save_case(
            {
                "case_id": "DJ-ANALYTICS2",
                "status": "completed",
                "created_at": "2026-07-28T08:00:00+00:00",
                "customer": {"customer_name": "已完成客户", "business_type": "TKP"},
                "trace": [
                    {"ts": "2026-07-28T08:00:00+00:00", "stage": "credit.approval_requested"},
                    {"ts": "2026-07-28T20:00:00+00:00", "stage": "credit.effective"},
                    {"ts": "2026-07-28T20:00:00+00:00", "stage": "workflow.interrupt"},
                    {"ts": "2026-07-29T08:00:00+00:00", "stage": "workflow.resumed"},
                    {"ts": "2026-07-29T10:00:00+00:00", "stage": "workflow.finalized"},
                ],
            }
        )
        self.save_case(
            {
                "case_id": "DJ-OLDCASE01",
                "status": "completed",
                "created_at": "2026-01-01T08:00:00+00:00",
                "customer": {"customer_name": "范围外客户", "business_type": "TKP"},
                "trace": [{"ts": "2026-01-02T08:00:00+00:00", "stage": "workflow.finalized"}],
            }
        )
        report = AnalyticsService(
            policy_path=str(self.policy_path), case_root=str(self.case_root)
        ).report(days=7, now=datetime(2026, 7, 31, tzinfo=timezone.utc))
        self.assertEqual(report["metrics"]["case_total"], 1)
        self.assertEqual(report["metrics"]["completed_cases"], 1)
        self.assertEqual(report["metrics"]["completion_rate"], 1.0)
        self.assertEqual(report["metrics"]["avg_cycle_hours"], 26.0)
        self.assertEqual(report["metrics"]["completed_tasks"], 2)
        self.assertEqual(report["metrics"]["sla_on_time_rate"], 1.0)
        self.assertEqual(sum(row["created"] for row in report["trend"]), 1)
        self.assertEqual(sum(row["completed"] for row in report["trend"]), 1)

    def test_csv_and_excel_exports_are_structured_and_readable(self):
        service = AnalyticsService(
            policy_path=str(self.policy_path), case_root=str(self.case_root)
        )
        report = service.report(days=30, now=datetime(2026, 7, 31, tzinfo=timezone.utc))
        csv_bytes = service.export_csv(report)
        self.assertTrue(csv_bytes.startswith(b"\xef\xbb\xbf"))
        csv_rows = list(csv.reader(io.StringIO(csv_bytes.decode("utf-8-sig"))))
        self.assertEqual(csv_rows[0][0], "案件号")

        excel_bytes = service.export_xlsx(report)
        workbook = load_workbook(io.BytesIO(excel_bytes), data_only=False)
        self.assertEqual(workbook.sheetnames, ["管理摘要", "节点明细", "案件明细"])
        self.assertIn("审批时效管理报表", workbook["管理摘要"]["A1"].value)
        self.assertEqual(workbook["节点明细"]["A4"].value, "案件号")
        self.assertEqual(workbook["案件明细"]["A4"].value, "案件号")
        self.assertEqual(len(workbook["管理摘要"]._charts), 1)

    def test_invalid_range_is_rejected(self):
        service = AnalyticsService(
            policy_path=str(self.policy_path), case_root=str(self.case_root)
        )
        with self.assertRaises(ValueError):
            service.report(days=14)


if __name__ == "__main__":
    unittest.main()
