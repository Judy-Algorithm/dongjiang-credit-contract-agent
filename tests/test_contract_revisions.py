import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

from dongjiang_agent.contract import ContractRevisionStore


CONTRACT = """销售合同

甲方：东江集团；乙方：修订测试客户。
合同标的：精密组件。合同金额：100万元。信用额度：100万元。
付款及账期：月结60天。知识产权：各自所有。保密：不得披露。
买方可随时取消订单且不承担任何责任。
违约责任：赔偿直接损失。解除与终止：违约催告后解除。争议解决：深圳法院。
"""


class ContractRevisionTests(unittest.TestCase):
    def test_text_revision_generates_clean_and_tracked_word_versions(self):
        previous = os.getcwd()
        with tempfile.TemporaryDirectory() as tmp:
            os.chdir(tmp)
            try:
                source = Path("data/archive/DJ-REVISION1/contract/source.txt")
                source.parent.mkdir(parents=True)
                source.write_text(CONTRACT, encoding="utf-8")
                case = {
                    "case_id": "DJ-REVISION1",
                    "source_documents": [
                        {
                            "document_id": "DOC-REVISION1",
                            "document_kind": "contract",
                            "name": "source.txt",
                            "sha256": "source-hash",
                            "archived_path": str(source.resolve()),
                            "media_type": "txt",
                        }
                    ],
                    "contract_reviews": [
                        {
                            "findings": [
                                {
                                    "rule_id": "DJ-CANCEL-WITHOUT-LIABILITY",
                                    "title": "客户可无责任取消订单或预测",
                                    "document_id": "DOC-REVISION1",
                                    "fragment_id": "line-6",
                                    "location": {"kind": "line", "line": 6},
                                }
                            ]
                        }
                    ],
                }
                before = source.read_bytes()
                revision = ContractRevisionStore().create(
                    case,
                    document_id="DOC-REVISION1",
                    decisions=[
                        {
                            "finding_key": "DJ-CANCEL-WITHOUT-LIABILITY:DOC-REVISION1:line-6",
                            "action": "accept",
                            "reason": "采用标准取消补偿机制",
                        }
                    ],
                    actor={"actor_id": "USR-SALES", "display_name": "销售甲"},
                )
                self.assertEqual(source.read_bytes(), before)
                self.assertEqual(revision["status"], "draft")
                manifest = ContractRevisionStore().get(
                    "DJ-REVISION1", revision["revision_id"]
                )
                clean = Path(manifest["artifacts"]["clean"]["path"])
                redline = Path(manifest["artifacts"]["redline"]["path"])
                clean_text = "\n".join(p.text for p in Document(clean).paragraphs)
                self.assertIn("至少提前30日书面通知", clean_text)
                self.assertNotIn("不承担任何责任", clean_text)
                with zipfile.ZipFile(redline) as archive:
                    xml = archive.read("word/document.xml").decode("utf-8")
                    settings = archive.read("word/settings.xml").decode("utf-8")
                self.assertIn("<w:del", xml)
                self.assertIn("<w:ins", xml)
                self.assertIn("不承担任何责任", xml)
                self.assertIn("至少提前30日书面通知", xml)
                self.assertIn("trackRevisions", settings)
                public_json = json.dumps(revision, ensure_ascii=False)
                self.assertNotIn(str(Path(tmp).resolve()), public_json)
                self.assertNotIn("不承担任何责任", public_json)
                clean_document = Document(clean)
                section = clean_document.sections[0]
                self.assertAlmostEqual(section.page_width.inches, 8.5, places=1)
                self.assertAlmostEqual(section.page_height.inches, 11.0, places=1)
                self.assertAlmostEqual(section.left_margin.inches, 1.0, places=1)
            finally:
                os.chdir(previous)

    def test_retain_requires_reason(self):
        previous = os.getcwd()
        with tempfile.TemporaryDirectory() as tmp:
            os.chdir(tmp)
            try:
                source = Path("data/archive/DJ-REVISION2/contract/source.txt")
                source.parent.mkdir(parents=True)
                source.write_text(CONTRACT, encoding="utf-8")
                case = {
                    "case_id": "DJ-REVISION2",
                    "source_documents": [{
                        "document_id": "DOC-REVISION2", "document_kind": "contract",
                        "name": "source.txt", "archived_path": str(source.resolve()), "media_type": "txt",
                    }],
                    "contract_reviews": [{"findings": [{
                        "rule_id": "DJ-CANCEL-WITHOUT-LIABILITY", "title": "取消风险",
                        "document_id": "DOC-REVISION2", "fragment_id": "line-6",
                        "location": {"kind": "line", "line": 6},
                    }]}],
                }
                with self.assertRaisesRegex(ValueError, "必须填写原因"):
                    ContractRevisionStore().create(
                        case,
                        document_id="DOC-REVISION2",
                        decisions=[{
                            "finding_key": "DJ-CANCEL-WITHOUT-LIABILITY:DOC-REVISION2:line-6",
                            "action": "retain", "reason": "",
                        }],
                        actor={"actor_id": "USR-LEGAL", "display_name": "法务甲"},
                    )
            finally:
                os.chdir(previous)

    def test_docx_table_cell_revision_preserves_table_and_tracks_change(self):
        previous = os.getcwd()
        with tempfile.TemporaryDirectory() as tmp:
            os.chdir(tmp)
            try:
                source = Path("data/archive/DJ-REVISION3/contract/source.docx")
                source.parent.mkdir(parents=True)
                document = Document()
                document.add_paragraph("销售合同")
                table = document.add_table(rows=2, cols=2)
                table.style = "Table Grid"
                table.cell(0, 0).text = "风险条款"
                table.cell(0, 1).text = "处置要求"
                table.cell(1, 0).text = "买方可随时取消订单且不承担任何责任。"
                table.cell(1, 1).text = "需要修改"
                document.save(source)
                case = {
                    "case_id": "DJ-REVISION3",
                    "source_documents": [{
                        "document_id": "DOC-REVISION3",
                        "document_kind": "contract",
                        "name": "source.docx",
                        "archived_path": str(source.resolve()),
                        "media_type": "docx",
                    }],
                    "contract_reviews": [{"findings": [{
                        "rule_id": "DJ-CANCEL-WITHOUT-LIABILITY",
                        "title": "取消风险",
                        "document_id": "DOC-REVISION3",
                        "fragment_id": "table-1-row-2-column-1",
                        "location": {
                            "kind": "word_table_cell", "table": 1,
                            "row": 2, "column": 1,
                        },
                    }]}],
                }
                revision = ContractRevisionStore().create(
                    case,
                    document_id="DOC-REVISION3",
                    decisions=[{
                        "finding_key": "DJ-CANCEL-WITHOUT-LIABILITY:DOC-REVISION3:table-1-row-2-column-1",
                        "action": "accept",
                        "reason": "采用标准取消补偿机制",
                    }],
                    actor={"actor_id": "USR-LEGAL", "display_name": "法务甲"},
                )
                manifest = ContractRevisionStore().get(
                    "DJ-REVISION3", revision["revision_id"]
                )
                clean = Document(manifest["artifacts"]["clean"]["path"])
                self.assertEqual(len(clean.tables), 1)
                self.assertIn("至少提前30日书面通知", clean.tables[0].cell(1, 0).text)
                self.assertEqual(clean.tables[0].cell(1, 1).text, "需要修改")
                with zipfile.ZipFile(manifest["artifacts"]["redline"]["path"]) as archive:
                    xml = archive.read("word/document.xml").decode("utf-8")
                self.assertIn("<w:tbl", xml)
                self.assertIn("<w:del", xml)
                self.assertIn("<w:ins", xml)
                self.assertIn("不承担任何责任", xml)
            finally:
                os.chdir(previous)


if __name__ == "__main__":
    unittest.main()
