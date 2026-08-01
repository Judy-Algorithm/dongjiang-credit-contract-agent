import json
import os
import tempfile
import unittest
from pathlib import Path

from docx import Document

from dongjiang_agent.contract import ContractTranslationStore, ContractTranslator
from dongjiang_agent.security import RedactionVault


class FakeGateway:
    available = True
    model = "translation-test-model"

    def __init__(self, *, mutate_token=False):
        self.mutate_token = mutate_token
        self.calls = []

    def analyze_redacted(self, redacted_text, instruction):
        self.calls.append((redacted_text, instruction))
        payload = json.loads(redacted_text.split("\n", 1)[1])
        translations = []
        for item in payload["segments"]:
            text = f"Translated: {item['text']}"
            if self.mutate_token:
                text = text.replace("⟦PARTY_", "⟦BROKEN_")
            translations.append(
                {"fragment_id": item["fragment_id"], "text": text}
            )
        return json.dumps({"translations": translations}, ensure_ascii=False)


class ContractTranslationTests(unittest.TestCase):
    def setUp(self):
        self.previous = os.getcwd()
        self.temp = tempfile.TemporaryDirectory()
        os.chdir(self.temp.name)

    def tearDown(self):
        os.chdir(self.previous)
        self.temp.cleanup()

    def _case(self):
        source = Path("data/archive/DJ-TRANS1/contract/source.docx")
        source.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        document.add_paragraph("销售合同")
        table = document.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        table.cell(0, 0).text = "合同主体"
        table.cell(0, 1).text = "付款条件"
        table.cell(1, 0).text = "甲方：测试公司"
        table.cell(1, 1).text = "月结60天"
        document.save(source)

        vault = RedactionVault("DJ-TRANS1", "data/vault")
        party = vault.redact("甲方：测试公司")
        vault.persist_local()
        return {
            "case_id": "DJ-TRANS1",
            "source_documents": [
                {
                    "document_id": "DOC-TRANS1",
                    "document_kind": "contract",
                    "name": "source.docx",
                    "archived_path": str(source.resolve()),
                    "sha256": "source-sha",
                    "media_type": "docx",
                    "parse_status": "parsed",
                    "fragments": [
                        {
                            "fragment_id": "paragraph-1",
                            "text": "销售合同",
                            "location": {"kind": "paragraph", "paragraph": 1},
                        },
                        {
                            "fragment_id": "table-1-row-1-column-1",
                            "text": "合同主体",
                            "location": {
                                "kind": "word_table_cell",
                                "table": 1,
                                "row": 1,
                                "column": 1,
                            },
                        },
                        {
                            "fragment_id": "table-1-row-1-column-2",
                            "text": "付款条件",
                            "location": {
                                "kind": "word_table_cell",
                                "table": 1,
                                "row": 1,
                                "column": 2,
                            },
                        },
                        {
                            "fragment_id": "table-1-row-2-column-1",
                            "text": party,
                            "location": {
                                "kind": "word_table_cell",
                                "table": 1,
                                "row": 2,
                                "column": 1,
                            },
                        },
                        {
                            "fragment_id": "table-1-row-2-column-2",
                            "text": "月结60天",
                            "location": {
                                "kind": "word_table_cell",
                                "table": 1,
                                "row": 2,
                                "column": 2,
                            },
                        },
                    ],
                }
            ],
        }

    def test_translation_draft_is_aligned_and_confirmed_export_preserves_table(self):
        gateway = FakeGateway()
        store = ContractTranslationStore(
            translator=ContractTranslator(gateway),
        )
        draft = store.create(
            self._case(),
            document_id="DOC-TRANS1",
            target_language="en",
            actor={"actor_id": "USR-LEGAL", "display_name": "法务甲"},
        )
        self.assertEqual(draft["status"], "draft")
        self.assertEqual(draft["segment_count"], 5)
        self.assertEqual(draft["entries"][3]["source_text"], "甲方：测试公司")
        self.assertNotIn("测试公司", gateway.calls[0][0])
        manifest = store.get("DJ-TRANS1", draft["translation_id"])
        self.assertNotIn(
            "测试公司", json.dumps(manifest, ensure_ascii=False)
        )

        confirmed = store.confirm(
            "DJ-TRANS1",
            draft["translation_id"],
            entries=[
                {
                    "fragment_id": item["fragment_id"],
                    "translated_text": item["translated_text"],
                }
                for item in draft["entries"]
            ],
            review_note="已逐段核对主体、付款条件和表格结构。",
            actor={"actor_id": "USR-LEGAL2", "display_name": "法务乙"},
        )
        self.assertEqual(confirmed["status"], "confirmed")
        target, _ = store.artifact_path("DJ-TRANS1", draft["translation_id"])
        exported = Document(target)
        self.assertEqual(len(exported.tables), 1)
        self.assertIn("机器翻译并经人工确认", exported.paragraphs[0].text)
        self.assertIn("Translated: 销售合同", exported.paragraphs[2].text)
        self.assertIn(
            "[英文译文] Translated: 月结60天",
            exported.tables[0].cell(1, 1).text,
        )
        self.assertIn("测试公司", exported.tables[0].cell(1, 0).text)
        public = json.dumps(store.list("DJ-TRANS1"), ensure_ascii=False)
        self.assertNotIn(str(Path(self.temp.name).resolve()), public)
        self.assertNotIn("测试公司", public)

    def test_translation_rejects_changed_redaction_token(self):
        store = ContractTranslationStore(
            translator=ContractTranslator(FakeGateway(mutate_token=True)),
        )
        with self.assertRaisesRegex(ValueError, "脱敏令牌"):
            store.create(
                self._case(),
                document_id="DOC-TRANS1",
                target_language="en",
                actor={"actor_id": "USR-LEGAL", "display_name": "法务甲"},
            )

    def test_translation_requires_human_review_note(self):
        store = ContractTranslationStore(
            translator=ContractTranslator(FakeGateway()),
        )
        draft = store.create(
            self._case(),
            document_id="DOC-TRANS1",
            target_language="ja",
            actor={"actor_id": "USR-LEGAL", "display_name": "法务甲"},
        )
        with self.assertRaisesRegex(ValueError, "人工复核说明"):
            store.confirm(
                "DJ-TRANS1",
                draft["translation_id"],
                entries=[
                    {
                        "fragment_id": item["fragment_id"],
                        "translated_text": item["translated_text"],
                    }
                    for item in draft["entries"]
                ],
                review_note="",
                actor={"actor_id": "USR-LEGAL", "display_name": "法务甲"},
            )

    @staticmethod
    def _confirmed_entries(draft):
        return [
            {
                "fragment_id": item["fragment_id"],
                "translations": dict(item["translations"]),
            }
            for item in draft["entries"]
        ]

    def test_multilingual_export_follows_requested_language_order(self):
        store = ContractTranslationStore(
            translator=ContractTranslator(FakeGateway()),
        )
        draft = store.create(
            self._case(),
            document_id="DOC-TRANS1",
            target_languages=["ja", "en", "vi"],
            actor={"actor_id": "USR-LEGAL", "display_name": "法务甲"},
        )
        self.assertEqual(draft["target_languages"], ["ja", "en", "vi"])
        self.assertEqual(draft["target_language_label"], "日语 / 英文 / 越南语")
        confirmed = store.confirm(
            "DJ-TRANS1",
            draft["translation_id"],
            entries=self._confirmed_entries(draft),
            review_note="已按日、英、越顺序核对。",
            actor={"actor_id": "USR-LEGAL2", "display_name": "法务乙"},
        )
        self.assertEqual(confirmed["status"], "confirmed")
        target, _ = store.artifact_path("DJ-TRANS1", draft["translation_id"])
        exported = Document(target)
        paragraphs = [item.text for item in exported.paragraphs]
        source_index = paragraphs.index("销售合同")
        self.assertTrue(paragraphs[source_index + 1].startswith("[日语译文]"))
        self.assertTrue(paragraphs[source_index + 2].startswith("[英文译文]"))
        self.assertTrue(paragraphs[source_index + 3].startswith("[越南语译文]"))

    def test_incremental_translation_only_sends_changed_fragment(self):
        gateway = FakeGateway()
        store = ContractTranslationStore(
            translator=ContractTranslator(gateway),
        )
        case = self._case()
        first = store.create(
            case,
            document_id="DOC-TRANS1",
            target_languages=["en", "ja"],
            actor={"actor_id": "USR-LEGAL", "display_name": "法务甲"},
        )
        store.confirm(
            "DJ-TRANS1",
            first["translation_id"],
            entries=self._confirmed_entries(first),
            review_note="首版译文已确认。",
            actor={"actor_id": "USR-LEGAL2", "display_name": "法务乙"},
        )
        gateway.calls.clear()
        changed = self._case()
        changed["source_documents"][0]["fragments"][-1]["text"] = "月结90天"
        second = store.create(
            changed,
            document_id="DOC-TRANS1",
            target_languages=["en", "ja"],
            actor={"actor_id": "USR-LEGAL", "display_name": "法务甲"},
        )
        self.assertEqual(second["base_translation_id"], first["translation_id"])
        self.assertEqual(second["reused_translation_count"], 8)
        self.assertEqual(second["translated_translation_count"], 2)
        self.assertEqual(len(gateway.calls), 2)
        for redacted_text, _instruction in gateway.calls:
            payload = json.loads(redacted_text.split("\n", 1)[1])
            self.assertEqual(len(payload["segments"]), 1)
            self.assertEqual(
                payload["segments"][0]["fragment_id"],
                "table-1-row-2-column-2",
            )
        changed_entry = second["entries"][-1]
        self.assertEqual(
            changed_entry["translation_sources"],
            {"en": "translated", "ja": "translated"},
        )


if __name__ == "__main__":
    unittest.main()
