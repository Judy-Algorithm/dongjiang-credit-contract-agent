import csv
import io
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from dongjiang_agent.operations import BenchmarkService


class BenchmarkServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.service = BenchmarkService(report_root=self.root / "reports")

    def tearDown(self):
        self.temp.cleanup()

    def test_suite_is_versioned_and_case_keys_are_unique(self):
        suite = self.service.load_suite()
        keys = [item["case_key"] for item in suite["cases"]]

        self.assertEqual(suite["suite_id"], "dongjiang-competition-benchmark")
        self.assertEqual(suite["version"], "1.0.0")
        self.assertEqual(suite["policy_status"], "PROPOSED_NOT_OFFICIAL")
        self.assertEqual(len(keys), 10)
        self.assertEqual(len(keys), len(set(keys)))

        duplicate_path = self.root / "duplicate.json"
        duplicate_path.write_text(
            json.dumps(
                {
                    "suite_id": "duplicate",
                    "version": "1",
                    "cases": [{"case_key": "same"}, {"case_key": "same"}],
                }
            ),
            encoding="utf-8",
        )
        with self.assertRaisesRegex(ValueError, "唯一"):
            BenchmarkService(suite_path=duplicate_path).load_suite()

    def test_double_run_passes_in_isolation_and_exports_no_source_text(self):
        suite = self.service.load_suite()
        report = self.service.run(repeats=2)
        metrics = report["metrics"]

        self.assertEqual(report["status"], "passed")
        self.assertEqual(metrics["case_passed"], 10)
        self.assertEqual(metrics["case_total"], 10)
        self.assertEqual(metrics["check_passed"], metrics["check_total"])
        self.assertEqual(metrics["privacy_pass_rate"], 1.0)
        self.assertEqual(metrics["determinism_rate"], 1.0)
        self.assertEqual(report["runtime"]["isolation"], "temporary_workspace")
        self.assertEqual(report["runtime"]["external_ai"], "disabled")
        self.assertEqual(report["runtime"]["enterprise_integrations"], "disabled")
        self.assertEqual({item.name for item in self.root.iterdir()}, {"reports"})
        self.assertTrue((self.root / "reports" / "latest.json").is_file())

        sensitive_values = [
            str(value)
            for case in suite["cases"]
            for value in case.get("sensitive_values") or []
        ]
        source_contracts = list((suite.get("templates") or {}).values())
        report_text = json.dumps(report, ensure_ascii=False)
        csv_bytes = self.service.export_csv(report)
        csv_text = csv_bytes.decode("utf-8-sig")
        csv_rows = list(csv.reader(io.StringIO(csv_text)))
        self.assertEqual(csv_rows[0][0], "案例")

        workbook = load_workbook(io.BytesIO(self.service.export_xlsx(report)))
        workbook_text = "\n".join(
            str(cell.value or "")
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        )
        for forbidden in sensitive_values + source_contracts:
            self.assertNotIn(forbidden, report_text)
            self.assertNotIn(forbidden, csv_text)
            self.assertNotIn(forbidden, workbook_text)

    def test_repeat_bounds_lock_and_nondeterminism_detection(self):
        for repeats in (0, 4):
            with self.assertRaisesRegex(ValueError, "1至3"):
                self.service.run(repeats=repeats)

        self.assertTrue(BenchmarkService._run_lock.acquire(blocking=False))
        try:
            with self.assertRaisesRegex(RuntimeError, "正在运行"):
                self.service.run(repeats=1)
        finally:
            BenchmarkService._run_lock.release()

        merged = self.service._merge_iterations(
            {"case_key": "unstable", "name": "不稳定案例", "tags": []},
            [
                {
                    "duration_ms": 1,
                    "checks": [],
                    "signature": {"decision": "pass"},
                },
                {
                    "duration_ms": 2,
                    "checks": [],
                    "signature": {"decision": "block"},
                },
            ],
        )
        determinism = next(
            item for item in merged["checks"] if item["metric"] == "determinism"
        )
        self.assertFalse(determinism["passed"])
        self.assertFalse(merged["passed"])


if __name__ == "__main__":
    unittest.main()
