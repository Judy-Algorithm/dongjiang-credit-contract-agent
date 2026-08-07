"""Generate the contract-focused competition demo files."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


OUTPUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("系统完整测试资料")


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def create_chinese_word(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("精密组件销售合同（中文附件）")
    run.bold = True
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor(20, 45, 58)

    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(3.3)
    table.columns[1].width = Cm(12.5)
    rows = [
        ("甲方", "东江精密制造（深圳）有限公司"),
        ("乙方", "华南新能源汽车技术有限公司及其子公司华南动力系统有限公司"),
        ("联系人", "陈敏，电话 13800138000，邮箱 min.chen@example.com"),
        ("合同标的", "PBT-GF30 精密注塑组件，图纸编号：DJ-TKP-2026-001"),
        ("合同金额", "人民币 3,600,000 元（含税）"),
        ("授信额度", "人民币 3,000,000 元"),
        ("付款", "付款期限：月结 120 天；以双方确认的发票与验收记录为准。"),
        ("交付", "按采购订单交付，所有权和风险于客户验收时转移。"),
        ("违约责任", "违约方仅赔偿可证明的直接损失，累计责任不超过合同金额的 30%。"),
        ("知识产权", "双方背景知识产权各自所有；履约资料仅限本合同目的使用，不可转让。"),
        ("保密", "双方对价格、技术参数、客户资料及商业秘密承担保密义务。"),
        ("解除与终止", "一方重大违约且书面催告后 30 日未改正，守约方可解除。"),
        ("争议解决", "适用中国法律，争议提交深圳国际仲裁院仲裁。"),
        ("收款账户", "6222021234567890123（仅为测试数据）"),
    ]
    for index, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        _set_cell_text(cells[0], label, bold=True)
        _set_cell_text(cells[1], value)
        _shade(cells[0], "E8EEF0")
        if index % 2:
            _shade(cells[1], "F8FAFA")
    for row in table.rows:
        row.cells[0].width = Cm(3.3)
        row.cells[1].width = Cm(12.5)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    note_run = note.add_run("演示用途：本文件包含虚构敏感字段与超账期条款，用于验证脱敏和合同例外授权。")
    note_run.italic = True
    note_run.font.size = Pt(8.5)
    note_run.font.color.rgb = RGBColor(92, 103, 109)
    document.core_properties.title = "东江合同 Agent 中文演示合同"
    document.core_properties.author = "东江集团比赛项目"
    document.save(path)


def create_vietnamese_pdf(path: Path) -> None:
    font_candidates = [
        Path(r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\DejaVuSans.ttf"),
    ]
    font_path = next((item for item in font_candidates if item.is_file()), None)
    font_name = "Helvetica"
    if font_path:
        font_name = "DemoUnicode"
        pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "DemoTitle", parent=styles["Title"], fontName=font_name,
        fontSize=17, leading=22, textColor=colors.HexColor("#142D3A"),
        alignment=TA_CENTER, spaceAfter=10,
    )
    body = ParagraphStyle(
        "DemoBody", parent=styles["BodyText"], fontName=font_name,
        fontSize=9.5, leading=14, textColor=colors.HexColor("#26343B"),
    )
    small = ParagraphStyle(
        "DemoSmall", parent=body, fontSize=8, leading=11,
        textColor=colors.HexColor("#5C676D"),
    )
    rows = [
        ("Bên mua", "Công ty TNHH Sao Việt Mobility và công ty con Sao Việt Components"),
        ("Bên bán", "Dongjiang Precision Manufacturing Co., Ltd."),
        ("Liên hệ", "Nguyễn An, +84 912345678, nguyen.an@example.vn"),
        ("Hàng hóa", "Linh kiện nhựa kỹ thuật; thông số kỹ thuật: PA66-GF35; mã bản vẽ: DJ-VN-2026-09"),
        ("Giá trị hợp đồng", "VND 8,000,000,000"),
        ("Hạn mức tín dụng", "VND 6,000,000,000"),
        ("Thanh toán", "Thời hạn thanh toán: 90 ngày kể từ ngày nghiệm thu và nhận hóa đơn hợp lệ."),
        ("Trách nhiệm", "Bên vi phạm chỉ bồi thường thiệt hại trực tiếp có chứng từ; tổng trách nhiệm không vượt quá 30% giá trị hợp đồng."),
        ("Sở hữu trí tuệ", "Quyền sở hữu trí tuệ nền tảng thuộc mỗi bên; giấy phép chỉ dùng để thực hiện hợp đồng, miễn phí và không được chuyển nhượng."),
        ("Bảo mật", "Các bên bảo mật giá, thông số kỹ thuật, dữ liệu khách hàng và bí mật kinh doanh."),
        ("Chấm dứt", "Một bên có thể chấm dứt sau khi thông báo vi phạm và cho 30 ngày để khắc phục."),
        ("Giải quyết tranh chấp", "Luật áp dụng là pháp luật Việt Nam; tranh chấp được giải quyết tại trọng tài VIAC."),
        ("Tài khoản", "9704360123456789 (dữ liệu thử nghiệm)"),
    ]
    story = [
        Paragraph("HỢP ĐỒNG MUA BÁN LINH KIỆN (PHỤ LỤC TIẾNG VIỆT)", title),
        Paragraph("Tài liệu trình diễn nhận diện ngôn ngữ, dịch thuật, che dữ liệu và phân tích rủi ro.", small),
        Spacer(1, 5 * mm),
    ]
    data = [[Paragraph("Mục", body), Paragraph("Nội dung", body)]] + [
        [Paragraph(label, body), Paragraph(value, body)] for label, value in rows
    ]
    table = Table(data, colWidths=[42 * mm, 130 * mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173D4D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 1), (0, -1), colors.HexColor("#E8EEF0")),
        ("ROWBACKGROUNDS", (1, 1), (1, -1), [colors.white, colors.HexColor("#F8FAFA")]),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C9D2D6")),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(table)
    story.extend([
        Spacer(1, 4 * mm),
        Paragraph("Lưu ý: Tên, số liên hệ, số tài khoản và số tiền trong tài liệu này đều là dữ liệu giả phục vụ kiểm thử.", small),
    ])
    document = SimpleDocTemplate(
        str(path), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm,
        topMargin=16 * mm, bottomMargin=16 * mm,
        title="Dongjiang Vietnamese Contract Demo",
        author="Dongjiang Competition Project",
    )
    document.build(story)


def create_bottom_line_word(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    title = document.add_paragraph("销售合同（底线条款退回案例）")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].bold = True
    title.runs[0].font.size = Pt(16)
    clauses = [
        "甲方：华南整车采购有限公司；乙方：东江精密制造（深圳）有限公司。",
        "合同标的：汽车精密注塑组件；合同金额：人民币 1,800,000 元；授信额度：人民币 1,200,000 元。",
        "付款：月结 60 天；知识产权由双方各自所有；双方承担保密义务。",
        "违约责任：违约方赔偿可证明的直接损失。",
        "底线条款：买方可随时取消采购订单且不承担任何责任、费用或赔偿，卖方应自行承担已采购物料及在制品损失。",
        "解除与终止：重大违约经 30 日催告后可解除。争议解决：适用中国法律，提交深圳法院管辖。",
    ]
    for index, clause in enumerate(clauses, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(7)
        run = paragraph.add_run(f"第 {index} 条  {clause}")
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)
        if index == 5:
            run.bold = True
            run.font.color.rgb = RGBColor(168, 46, 35)
    document.add_paragraph("本文件仅用于验证命中底线条款后退回业务修改。")
    document.save(path)


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    create_chinese_word(OUTPUT / "12-四角色主案例-中文合同.docx")
    create_vietnamese_pdf(OUTPUT / "14-四角色主案例-越南语合同.pdf")
    create_bottom_line_word(OUTPUT / "15-底线条款退回案例-中文合同.docx")
    print(OUTPUT.resolve())


if __name__ == "__main__":
    main()
