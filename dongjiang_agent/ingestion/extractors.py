"""Multi-format extraction with zero-dependency fallbacks."""

from __future__ import annotations

import csv
import html
import io
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any
from xml.etree import ElementTree


@dataclass(slots=True)
class DocumentFragment:
    fragment_id: str
    text: str
    location: dict[str, Any]


@dataclass(slots=True)
class ExtractedDocument:
    path: str
    media_type: str
    text: str
    extractor: str
    warnings: list[str]
    fragments: list[DocumentFragment]


class DocumentExtractor:
    TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm"}
    IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}
    MAX_OOXML_MEMBERS = 10_000
    MAX_OOXML_MEMBER_BYTES = 64 * 1024 * 1024
    MAX_OOXML_UNCOMPRESSED_BYTES = 128 * 1024 * 1024

    def extract(self, path: str | Path) -> ExtractedDocument:
        target = Path(path).expanduser().resolve()
        if not target.is_file():
            raise FileNotFoundError(f"文件不存在：{target}")
        suffix = target.suffix.lower()
        if suffix in self.TEXT_EXTENSIONS:
            text = self._read_text(target)
            return ExtractedDocument(
                str(target),
                suffix.lstrip("."),
                text,
                "native-text",
                [],
                self._line_fragments(text),
            )
        if suffix == ".docx":
            self._validate_ooxml_archive(target)
            text, fragments = self._docx_content(target)
            return ExtractedDocument(str(target), "docx", text, "ooxml", [], fragments)
        if suffix == ".xlsx":
            self._validate_ooxml_archive(target)
            text, fragments = self._xlsx_content(target)
            return ExtractedDocument(str(target), "xlsx", text, "ooxml", [], fragments)
        if suffix == ".pdf":
            text, extractor, warnings, fragments = self._pdf_content(target)
            return ExtractedDocument(str(target), "pdf", text, extractor, warnings, fragments)
        if suffix in self.IMAGE_EXTENSIONS:
            text, confidence, warnings, regions = self._image_text(target)
            fragments = [
                DocumentFragment(
                    "image-1",
                    text,
                    {
                        "kind": "image",
                        "image": 1,
                        "ocr": True,
                        "ocr_confidence": confidence,
                        "ocr_regions": regions,
                    },
                )
            ] if text.strip() else []
            return ExtractedDocument(
                str(target), suffix.lstrip("."), text, "tesseract", warnings, fragments
            )
        raise ValueError(f"暂不支持的文件格式：{suffix or '无扩展名'}")

    @classmethod
    def _validate_ooxml_archive(
        cls,
        path: Path,
        *,
        max_members: int | None = None,
        max_member_bytes: int | None = None,
        max_uncompressed_bytes: int | None = None,
    ) -> None:
        """Reject malformed or expansion-heavy DOCX/XLSX files before parsing."""
        member_limit = max_members or cls.MAX_OOXML_MEMBERS
        member_size_limit = max_member_bytes or cls.MAX_OOXML_MEMBER_BYTES
        total_limit = max_uncompressed_bytes or cls.MAX_OOXML_UNCOMPRESSED_BYTES
        try:
            with zipfile.ZipFile(path) as archive:
                members = archive.infolist()
        except (OSError, zipfile.BadZipFile) as exc:
            raise ValueError("DOCX/XLSX文件损坏或不是有效的OOXML文档。") from exc
        if len(members) > member_limit:
            raise ValueError("DOCX/XLSX内部文件过多，已拒绝解析。")
        total = 0
        for member in members:
            if member.file_size < 0 or member.file_size > member_size_limit:
                raise ValueError("DOCX/XLSX内部单个文件过大，已拒绝解析。")
            total += member.file_size
            if total > total_limit:
                raise ValueError("DOCX/XLSX解压后内容过大，已拒绝解析。")

    @staticmethod
    def _read_text(path: Path) -> str:
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                text = raw.decode(encoding)
                if path.suffix.lower() == ".csv":
                    rows = csv.reader(io.StringIO(text))
                    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
                if path.suffix.lower() in {".html", ".htm"}:
                    text = re.sub(r"<[^>]+>", " ", text)
                    return html.unescape(re.sub(r"\s+", " ", text))
                return text
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", errors="replace")

    @staticmethod
    def _docx_content(path: Path) -> tuple[str, list[DocumentFragment]]:
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml")
        root = ElementTree.fromstring(xml)
        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        lines: list[str] = []
        fragments: list[DocumentFragment] = []
        table_paragraphs = {
            paragraph
            for table in root.findall(".//w:tbl", namespace)
            for paragraph in table.findall(".//w:p", namespace)
        }
        paragraph_number = 0
        for paragraph in root.findall(".//w:p", namespace):
            text = DocumentExtractor._word_paragraph_text(paragraph, namespace)
            if not text:
                continue
            paragraph_number += 1
            if paragraph not in table_paragraphs:
                lines.append(text)
                fragments.append(
                    DocumentFragment(
                        f"paragraph-{paragraph_number}",
                        text,
                        {"kind": "paragraph", "paragraph": paragraph_number},
                    )
                )

        for table_number, table in enumerate(root.findall(".//w:tbl", namespace), start=1):
            table_lines: list[str] = []
            vertical_merges: dict[int, DocumentFragment] = {}
            for row_number, row in enumerate(table.findall("./w:tr", namespace), start=1):
                column = 1
                row_values: list[str] = []
                for cell in row.findall("./w:tc", namespace):
                    properties = cell.find("./w:tcPr", namespace)
                    grid_span_node = (
                        properties.find("./w:gridSpan", namespace)
                        if properties is not None
                        else None
                    )
                    try:
                        column_span = max(
                            1,
                            int(
                                grid_span_node.get(
                                    f"{{{namespace['w']}}}val", "1"
                                )
                                if grid_span_node is not None
                                else 1
                            ),
                        )
                    except (TypeError, ValueError):
                        column_span = 1
                    merge_node = (
                        properties.find("./w:vMerge", namespace)
                        if properties is not None
                        else None
                    )
                    merge_value = (
                        merge_node.get(f"{{{namespace['w']}}}val", "continue")
                        if merge_node is not None
                        else ""
                    )
                    cell_text = "\n".join(
                        value
                        for value in (
                            DocumentExtractor._word_paragraph_text(paragraph, namespace)
                            for paragraph in cell.findall("./w:p", namespace)
                        )
                        if value
                    ).strip()
                    if merge_node is not None and merge_value != "restart":
                        origin = vertical_merges.get(column)
                        if origin is not None:
                            origin.location["row_span"] = (
                                int(origin.location.get("row_span") or 1) + 1
                            )
                        column += column_span
                        continue
                    location = {
                        "kind": "word_table_cell",
                        "table": table_number,
                        "row": row_number,
                        "column": column,
                        "row_span": 1,
                        "column_span": column_span,
                    }
                    if cell_text:
                        fragment = DocumentFragment(
                            f"table-{table_number}-row-{row_number}-column-{column}",
                            cell_text,
                            location,
                        )
                        fragments.append(fragment)
                        row_values.append(cell_text.replace("\n", " "))
                        if merge_node is not None:
                            for merged_column in range(column, column + column_span):
                                vertical_merges[merged_column] = fragment
                    elif merge_node is None:
                        for merged_column in range(column, column + column_span):
                            vertical_merges.pop(merged_column, None)
                    column += column_span
                if row_values:
                    table_lines.append(" | ".join(row_values))
            if table_lines:
                lines.append(f"[Word表格 {table_number}]")
                lines.extend(table_lines)
        return "\n".join(lines), fragments

    @staticmethod
    def _word_paragraph_text(
        paragraph: ElementTree.Element,
        namespace: dict[str, str],
    ) -> str:
        return "".join(
            node.text or "" for node in paragraph.findall(".//w:t", namespace)
        ).strip()

    @staticmethod
    def _xlsx_content(path: Path) -> tuple[str, list[DocumentFragment]]:
        try:
            import openpyxl  # type: ignore
        except ImportError:
            return DocumentExtractor._xlsx_content_ooxml(path)
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        lines: list[str] = []
        fragments: list[DocumentFragment] = []
        for sheet in workbook.worksheets:
            lines.append(f"[工作表] {sheet.title}")
            for row in sheet.iter_rows():
                values = [str(cell.value).strip() if cell.value is not None else "" for cell in row]
                if any(values):
                    lines.append(" | ".join(values))
                for cell, value in zip(row, values):
                    if not value:
                        continue
                    fragments.append(
                        DocumentFragment(
                            f"cell-{sheet.title}-{cell.coordinate}",
                            value,
                            {
                                "kind": "cell",
                                "sheet": sheet.title,
                                "cell": cell.coordinate,
                            },
                        )
                    )
        workbook.close()
        return "\n".join(lines), fragments

    @staticmethod
    def _xlsx_content_ooxml(path: Path) -> tuple[str, list[DocumentFragment]]:
        with zipfile.ZipFile(path) as archive:
            shared: list[str] = []
            if "xl/sharedStrings.xml" in archive.namelist():
                root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
                shared = ["".join(node.itertext()) for node in root]
            lines: list[str] = []
            fragments: list[DocumentFragment] = []
            for name in sorted(item for item in archive.namelist() if re.match(r"xl/worksheets/sheet\d+\.xml", item)):
                sheet_name = Path(name).stem
                lines.append(f"[工作表] {sheet_name}")
                root = ElementTree.fromstring(archive.read(name))
                for row in root.iter():
                    if not row.tag.endswith("}row"):
                        continue
                    values: list[str] = []
                    for cell in row:
                        if not cell.tag.endswith("}c"):
                            continue
                        raw = next((node.text or "" for node in cell if node.tag.endswith("}v")), "")
                        if cell.attrib.get("t") == "s" and raw.isdigit() and int(raw) < len(shared):
                            raw = shared[int(raw)]
                        values.append(raw)
                        if raw:
                            coordinate = str(cell.attrib.get("r") or "")
                            fragments.append(
                                DocumentFragment(
                                    f"cell-{sheet_name}-{coordinate}",
                                    raw,
                                    {
                                        "kind": "cell",
                                        "sheet": sheet_name,
                                        "cell": coordinate,
                                    },
                                )
                            )
                    if any(values):
                        lines.append(" | ".join(values))
            return "\n".join(lines), fragments

    @classmethod
    def _pdf_content(
        cls,
        path: Path,
    ) -> tuple[str, str, list[str], list[DocumentFragment]]:
        warnings: list[str] = []
        try:
            from pypdf import PdfReader  # type: ignore
            reader = PdfReader(str(path))
            pages = [page.extract_text() or "" for page in reader.pages]
            empty_pages = [
                number
                for number, content in enumerate(pages, start=1)
                if not content.strip()
            ]
            ocr_pages: dict[int, tuple[Any, ...]] = {}
            if empty_pages:
                ocr_pages, ocr_warnings = cls._ocr_pdf_pages(path, empty_pages)
                warnings.extend(ocr_warnings)
                for number, result in ocr_pages.items():
                    content = str(result[0])
                    pages[number - 1] = content
            text = "\n".join(pages)
            if text.strip():
                fragments = [
                    DocumentFragment(
                        f"page-{number}",
                        content,
                        {
                            "kind": "page",
                            "page": number,
                            "ocr": number in ocr_pages,
                            "ocr_confidence": (
                                ocr_pages[number][1] if number in ocr_pages else None
                            ),
                            **(
                                {"ocr_regions": ocr_pages[number][2]}
                                if number in ocr_pages
                                and len(ocr_pages[number]) > 2
                                and ocr_pages[number][2]
                                else {}
                            ),
                        },
                    )
                    for number, content in enumerate(pages, start=1)
                    if content.strip()
                ]
                extractor = "pypdf+tesseract" if ocr_pages else "pypdf"
                return text, extractor, warnings, fragments
            if empty_pages and not warnings:
                warnings.append("扫描版 PDF 未识别出可用文本。")
        except ImportError:
            pass
        except Exception as exc:
            warnings.append(f"pypdf 解析失败：{type(exc).__name__}")
        binary = shutil.which("pdftotext")
        if binary:
            result = subprocess.run(
                [binary, "-layout", str(path), "-"],
                check=False,
                capture_output=True,
                text=True,
                timeout=60,
            )
            if result.stdout.strip():
                pages = result.stdout.split("\f")
                fragments = [
                    DocumentFragment(
                        f"page-{number}",
                        content,
                        {"kind": "page", "page": number},
                    )
                    for number, content in enumerate(pages, start=1)
                    if content.strip()
                ]
                return result.stdout, "pdftotext", warnings, fragments
        warnings.append("PDF 无可提取文本，请安装 pypdf、pypdfium2 和 Tesseract OCR。")
        return "", "none", warnings, []

    @classmethod
    def _ocr_pdf_pages(
        cls,
        path: Path,
        page_numbers: list[int],
    ) -> tuple[dict[int, tuple[Any, ...]], list[str]]:
        binary = cls._tesseract_binary()
        if not binary:
            return {}, ["扫描页需要 OCR，但当前环境未安装 Tesseract。"]
        results: dict[int, tuple[Any, ...]] = {}
        warnings: list[str] = []
        with tempfile.TemporaryDirectory(prefix="dongjiang-pdf-ocr-") as tmp:
            root = Path(tmp)
            for page_number in page_numbers:
                image_path = root / f"page-{page_number}.png"
                try:
                    cls._render_pdf_page(path, page_number, image_path)
                    text, confidence, page_warnings, regions = cls._tesseract_text(
                        image_path, binary=binary
                    )
                    warnings.extend(
                        f"第 {page_number} 页：{warning}" for warning in page_warnings
                    )
                    if text.strip():
                        results[page_number] = (text, confidence, regions)
                    else:
                        warnings.append(f"第 {page_number} 页 OCR 未识别出文本。")
                except Exception as exc:
                    warnings.append(
                        f"第 {page_number} 页 OCR 失败：{type(exc).__name__}"
                    )
        return results, warnings

    @staticmethod
    def _render_pdf_page(path: Path, page_number: int, target: Path) -> None:
        try:
            import pypdfium2 as pdfium  # type: ignore

            document = pdfium.PdfDocument(str(path))
            try:
                page = document[page_number - 1]
                try:
                    page.render(scale=3).to_pil().save(target, format="PNG")
                finally:
                    page.close()
            finally:
                document.close()
            return
        except ImportError:
            pass
        binary = shutil.which("pdftoppm")
        if not binary:
            raise RuntimeError("未安装 pypdfium2 或 pdftoppm，无法渲染扫描 PDF。")
        prefix = target.with_suffix("")
        result = subprocess.run(
            [
                binary,
                "-f",
                str(page_number),
                "-l",
                str(page_number),
                "-singlefile",
                "-r",
                "300",
                "-png",
                str(path),
                str(prefix),
            ],
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=120,
        )
        if result.returncode != 0 or not target.is_file():
            raise RuntimeError(result.stderr.strip() or "PDF 页面渲染失败。")

    @staticmethod
    def _line_fragments(text: str) -> list[DocumentFragment]:
        fragments: list[DocumentFragment] = []
        for number, line in enumerate(text.splitlines(), start=1):
            content = line.strip()
            if content:
                fragments.append(
                    DocumentFragment(
                        f"line-{number}",
                        content,
                        {"kind": "line", "line": number},
                    )
                )
        if not fragments and text.strip():
            fragments.append(
                DocumentFragment("line-1", text.strip(), {"kind": "line", "line": 1})
            )
        return fragments

    @classmethod
    def _image_text(
        cls, path: Path
    ) -> tuple[str, float | None, list[str], list[dict[str, Any]]]:
        binary = cls._tesseract_binary()
        if not binary:
            return "", None, ["未找到 Tesseract，图片已登记但未执行 OCR。"], []
        return cls._tesseract_text(path, binary=binary)

    @classmethod
    def _tesseract_text(
        cls,
        path: Path,
        *,
        binary: str,
    ) -> tuple[str, float | None, list[str], list[dict[str, Any]]]:
        environment, tessdata_prefix = cls._tesseract_environment()
        languages, language_warnings = cls._tesseract_languages(
            binary, tessdata_prefix
        )
        result = subprocess.run(
            [binary, str(path), "stdout", "-l", languages, "--psm", "6", "tsv"],
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            env=environment,
            timeout=120,
        )
        warnings = list(language_warnings)
        if result.returncode != 0:
            fallback = subprocess.run(
                [binary, str(path), "stdout", "-l", languages, "--psm", "6"],
                check=False,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                env=environment,
                timeout=120,
            )
            if fallback.returncode == 0 and fallback.stdout.strip():
                warnings.append("OCR 置信度不可用，已回退到纯文本识别。")
                return fallback.stdout.strip(), None, warnings, []
            warnings.append(
                fallback.stderr.strip()
                or result.stderr.strip()
                or "OCR 失败"
            )
            return "", None, warnings, []
        rows = csv.DictReader(io.StringIO(result.stdout), delimiter="\t")
        lines: dict[tuple[str, str, str, str], dict[str, Any]] = {}
        confidences: list[float] = []
        for row in rows:
            value = str(row.get("text") or "").strip()
            if not value:
                continue
            key = tuple(str(row.get(name) or "0") for name in ("page_num", "block_num", "par_num", "line_num"))
            line = lines.setdefault(
                key,
                {"words": [], "left": [], "top": [], "right": [], "bottom": []},
            )
            line["words"].append(value)
            try:
                left = int(row.get("left") or 0)
                top = int(row.get("top") or 0)
                width = int(row.get("width") or 0)
                height = int(row.get("height") or 0)
                line["left"].append(left)
                line["top"].append(top)
                line["right"].append(left + width)
                line["bottom"].append(top + height)
            except (TypeError, ValueError):
                pass
            try:
                confidence = float(row.get("conf") or -1)
                if confidence >= 0:
                    confidences.append(confidence)
            except (TypeError, ValueError):
                pass
        try:
            from PIL import Image  # type: ignore

            with Image.open(path) as image:
                image_width, image_height = image.size
        except Exception:
            image_width = max(
                (max(item["right"], default=1) for item in lines.values()),
                default=1,
            )
            image_height = max(
                (max(item["bottom"], default=1) for item in lines.values()),
                default=1,
            )
        normalized_lines: list[str] = []
        regions: list[dict[str, Any]] = []
        cursor = 0
        for line in lines.values():
            normalized = cls._normalize_ocr_text(" ".join(line["words"]))
            if not normalized:
                continue
            start = cursor
            cursor += len(re.sub(r"\s+", "", normalized))
            normalized_lines.append(normalized)
            if line["left"] and image_width and image_height:
                left = min(line["left"])
                top = min(line["top"])
                right = max(line["right"])
                bottom = max(line["bottom"])
                regions.append({
                    "start": start,
                    "end": cursor,
                    "x": round(left / image_width, 6),
                    "y": round(top / image_height, 6),
                    "width": round((right - left) / image_width, 6),
                    "height": round((bottom - top) / image_height, 6),
                })
        text = "\n".join(normalized_lines)
        confidence = (
            round(sum(confidences) / len(confidences) / 100, 3)
            if confidences
            else None
        )
        return text, confidence, warnings, regions

    @staticmethod
    def _normalize_ocr_text(text: str) -> str:
        cjk = r"\u3400-\u9fff\u3040-\u30ff"
        normalized = re.sub(
            rf"(?<=[{cjk}])[ \t]+(?=[{cjk}])",
            "",
            str(text or ""),
        )
        normalized = re.sub(r"[ \t]+([：:，,。；;])", r"\1", normalized)
        normalized = re.sub(
            rf"([：:，,。；;])[ \t]+(?=[{cjk}])", r"\1", normalized
        )
        return normalized.strip()

    @staticmethod
    def _tesseract_binary() -> str | None:
        configured = os.getenv("DONGJIANG_TESSERACT_CMD", "").strip()
        if configured and Path(configured).is_file():
            return configured
        discovered = shutil.which("tesseract")
        if discovered:
            return discovered
        if os.name == "nt":
            for candidate in (
                Path(os.getenv("ProgramFiles", "C:/Program Files"))
                / "Tesseract-OCR/tesseract.exe",
                Path(os.getenv("LOCALAPPDATA", ""))
                / "Programs/Tesseract-OCR/tesseract.exe",
            ):
                if candidate.is_file():
                    return str(candidate)
            try:
                import winreg

                roots = (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER)
                paths = (
                    r"Software\Microsoft\Windows\CurrentVersion\Uninstall",
                    r"Software\WOW6432Node\Microsoft\Windows\CurrentVersion\Uninstall",
                )
                for root in roots:
                    for registry_path in paths:
                        try:
                            with winreg.OpenKey(root, registry_path) as uninstall:
                                for index in range(winreg.QueryInfoKey(uninstall)[0]):
                                    key_name = winreg.EnumKey(uninstall, index)
                                    with winreg.OpenKey(uninstall, key_name) as package:
                                        try:
                                            display_name = str(
                                                winreg.QueryValueEx(package, "DisplayName")[0]
                                            )
                                            uninstall_string = str(
                                                winreg.QueryValueEx(package, "UninstallString")[0]
                                            )
                                        except OSError:
                                            continue
                                        if "tesseract" not in display_name.lower():
                                            continue
                                        candidate = (
                                            Path(uninstall_string.strip('"')).parent
                                            / "tesseract.exe"
                                        )
                                        if candidate.is_file():
                                            return str(candidate)
                        except OSError:
                            continue
            except (ImportError, OSError):
                pass
        return None

    @staticmethod
    def _tesseract_environment() -> tuple[dict[str, str], str]:
        environment = dict(os.environ)
        configured = os.getenv("DONGJIANG_TESSDATA_PREFIX", "").strip()
        local = Path("data/tessdata").resolve()
        prefix = configured or (str(local) if local.is_dir() else "")
        if prefix:
            environment["TESSDATA_PREFIX"] = prefix
        return environment, prefix

    @staticmethod
    @lru_cache(maxsize=8)
    def _tesseract_languages(
        binary: str,
        tessdata_prefix: str = "",
    ) -> tuple[str, tuple[str, ...]]:
        requested = [
            item.strip()
            for item in re.split(
                r"[+,]",
                os.getenv(
                    "DONGJIANG_OCR_LANGUAGES", "chi_sim+eng+vie+jpn+spa"
                ),
            )
            if item.strip()
        ]
        environment = dict(os.environ)
        if tessdata_prefix:
            environment["TESSDATA_PREFIX"] = tessdata_prefix
        result = subprocess.run(
            [binary, "--list-langs"],
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            env=environment,
            timeout=15,
        )
        available = {
            line.strip()
            for line in result.stdout.splitlines()
            if line.strip() and "available languages" not in line.lower()
        }
        selected = [item for item in requested if item in available]
        if not selected:
            selected = ["eng"] if "eng" in available else sorted(available)[:1]
        if not selected:
            return "eng", ("无法读取 Tesseract 语言包，已尝试使用英文。",)
        missing = [item for item in requested if item not in available]
        warnings = (
            (f"未安装 OCR 语言包：{', '.join(missing)}；已使用 {', '.join(selected)}。",)
            if missing
            else ()
        )
        return "+".join(selected), warnings
