"""Historical workflow analytics and management-report exports."""

from __future__ import annotations

import csv
import io
import math
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from typing import Any

from ..persistence import CaseRepository
from .sla import (
    OWNER_WAITING,
    WAITING_BY_STATUS,
    WAITING_ROLES,
    case_sla,
    load_sla_policy,
)


TERMINAL_STATUSES = {
    "completed",
    "approved",
    "approved_by_exception",
    "approved_after_manual_review",
    "rejected",
    "credit_rejected",
    "credit_control_rejected",
    "closed",
}
SUCCESS_STATUSES = {
    "completed",
    "approved",
    "approved_by_exception",
    "approved_after_manual_review",
}
EXIT_STAGES = {
    "credit_approval": {
        "credit.effective",
        "credit.supplement_requested",
        "credit.rejected",
    },
    "credit_supplement": {"credit.supplemented", "credit.closed"},
    "special_release": {
        "credit.special_release_approved",
        "credit.special_release_rejected",
    },
    "contract_upload": {"workflow.resumed", "workflow.closed"},
    "sales_revision": {"sales.revised", "sales.closed"},
    "manager_approval": {"manager.approved", "manager.rejected"},
    "finance_legal_review": {
        "contract.approved",
        "manual.approved",
        "manual.supplemented",
        "manual.revision_requested",
    },
}
ROLE_LABELS = {
    "credit_approver": "信用审批人",
    "legal_reviewer": "合同法务",
    "exception_approver": "授权审批人",
}
STATUS_LABELS = {
    "completed": "已完成",
    "approved": "已批准",
    "approved_by_exception": "特批通过",
    "approved_after_manual_review": "复核通过",
    "rejected": "已拒绝",
    "credit_rejected": "信用申请已拒绝",
    "credit_control_rejected": "特别放行已拒绝",
    "closed": "已关闭",
    "credit_pending_approval": "等待信用审批",
    "credit_supplement_required": "等待补充信用资料",
    "credit_control_locked": "信用控制已锁定",
    "awaiting_contract": "等待上传合同",
    "blocked": "等待修改合同",
    "pending_special_approval": "等待管理层审批",
    "pending_manual_review": "等待财务法务复核",
    "pending_contract_approval": "等待合同法务审批",
}


def _utc_now() -> datetime:
    return datetime.now(timezone.utc)


def _parse_datetime(value: object) -> datetime | None:
    raw = str(value or "").strip()
    if not raw:
        return None
    try:
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
    except ValueError:
        return None
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)


def _normalized_now(value: datetime | None) -> datetime:
    current = value or _utc_now()
    return current if current.tzinfo else current.replace(tzinfo=timezone.utc)


def _excel_datetime(value: object) -> datetime | None:
    parsed = _parse_datetime(value)
    return parsed.astimezone(timezone.utc).replace(tzinfo=None) if parsed else None


def _trace(case: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[tuple[datetime, int, dict[str, Any]]] = []
    for index, item in enumerate(case.get("trace") or []):
        occurred_at = _parse_datetime(item.get("ts"))
        if occurred_at:
            rows.append((occurred_at, index, dict(item)))
    rows.sort(key=lambda row: (row[0], row[1]))
    return [item for _, _, item in rows]


def _decision_task(item: dict[str, Any], case: dict[str, Any]) -> str:
    data = dict(item.get("data") or {})
    decision = str(data.get("decision") or case.get("decision") or "")
    if decision == "block":
        return "sales_revision"
    if decision == "special_approval":
        return "manager_approval"
    if decision in {"manual_review", "pass"}:
        return "finance_legal_review"
    return ""


def _special_release_entry(
    case: dict[str, Any], trace: list[dict[str, Any]], index: int
) -> bool:
    for item in trace[index + 1 :]:
        stage = str(item.get("stage") or "")
        if stage in EXIT_STAGES["special_release"]:
            return True
        if stage in {"workflow.interrupt", "contract.completed", "workflow.finalized"}:
            return False
    return str(case.get("status") or "") == "credit_control_locked"


def _entry_task(
    case: dict[str, Any], trace: list[dict[str, Any]], index: int
) -> str:
    item = trace[index]
    stage = str(item.get("stage") or "")
    if stage == "credit.approval_requested":
        return "credit_approval"
    if stage in {"credit.insufficient_data", "credit.supplement_requested"}:
        return "credit_supplement"
    if stage == "credit.effective" and _special_release_entry(case, trace, index):
        return "special_release"
    if stage in {"workflow.interrupt", "credit.special_release_approved"}:
        return "contract_upload"
    if stage == "decision.routed":
        return _decision_task(item, case)
    if stage in {"manager.rejected", "manual.revision_requested"}:
        return "sales_revision"
    return ""


def _responsible(case: dict[str, Any], task: str) -> dict[str, str]:
    if task in OWNER_WAITING:
        owner = dict(case.get("owner") or case.get("applicant") or {})
        return {
            "key": str(owner.get("user_id") or "unassigned-owner"),
            "label": str(owner.get("display_name") or owner.get("username") or "未分配负责人"),
            "kind": "user",
        }
    roles = WAITING_ROLES.get(task) or []
    label = "/".join(ROLE_LABELS.get(role, role) for role in roles) or "按角色分配"
    return {"key": "roles:" + ",".join(roles), "label": label, "kind": "role"}


def _interval(
    case: dict[str, Any],
    task: str,
    entered_at: datetime,
    exited_at: datetime | None,
    target_hours: float,
    current: datetime,
) -> dict[str, Any]:
    end = exited_at or current
    duration = max((end - entered_at).total_seconds() / 3600, 0.0)
    completed = exited_at is not None
    customer = dict(case.get("customer") or {})
    responsible = _responsible(case, task)
    return {
        "case_id": str(case.get("case_id") or ""),
        "customer_name": str(customer.get("customer_name") or ""),
        "business_type": str(customer.get("business_type") or ""),
        "task": task,
        "entered_at": entered_at.isoformat(timespec="seconds"),
        "exited_at": exited_at.isoformat(timespec="seconds") if exited_at else "",
        "duration_hours": round(duration, 1),
        "target_hours": target_hours,
        "completed": completed,
        "met_sla": duration <= target_hours,
        "responsible": responsible,
    }


def task_intervals(
    case: dict[str, Any],
    *,
    now: datetime | None = None,
    policy: dict[str, Any] | None = None,
) -> list[dict[str, Any]]:
    """Return completed and current SLA-node intervals reconstructed from trace."""
    current = _normalized_now(now)
    rules = (policy or load_sla_policy()).get("tasks") or {}
    trace = _trace(case)
    open_tasks: dict[str, datetime] = {}
    intervals: list[dict[str, Any]] = []
    for index, item in enumerate(trace):
        stage = str(item.get("stage") or "")
        occurred_at = _parse_datetime(item.get("ts"))
        if not occurred_at:
            continue
        for task, exit_stages in EXIT_STAGES.items():
            entered_at = open_tasks.get(task)
            if entered_at and stage in exit_stages:
                rule = dict(rules.get(task) or {})
                intervals.append(
                    _interval(
                        case,
                        task,
                        entered_at,
                        occurred_at,
                        max(float(rule.get("target_hours") or 1), 1.0),
                        current,
                    )
                )
                del open_tasks[task]
        task = _entry_task(case, trace, index)
        if task and task in rules and task not in open_tasks:
            open_tasks[task] = occurred_at

    current_task = WAITING_BY_STATUS.get(str(case.get("status") or ""))
    if current_task and current_task in rules:
        sla = case_sla(case, now=current, policy=policy or load_sla_policy())
        entered_at = _parse_datetime(sla.get("entered_at"))
        if entered_at:
            open_tasks[current_task] = entered_at
    for task, entered_at in open_tasks.items():
        if task != current_task:
            continue
        rule = dict(rules.get(task) or {})
        intervals.append(
            _interval(
                case,
                task,
                entered_at,
                None,
                max(float(rule.get("target_hours") or 1), 1.0),
                current,
            )
        )
    intervals.sort(key=lambda item: (item["entered_at"], item["task"]))
    return intervals


def _percentile(values: list[float], percentile: float) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    index = max(math.ceil(len(ordered) * percentile) - 1, 0)
    return round(ordered[index], 1)


def _average(values: list[float]) -> float | None:
    return round(sum(values) / len(values), 1) if values else None


def _terminal_at(case: dict[str, Any]) -> datetime | None:
    if str(case.get("status") or "") not in TERMINAL_STATUSES:
        return None
    trace = _trace(case)
    terminal_stages = {
        "workflow.finalized",
        "workflow.closed",
        "sales.closed",
        "credit.closed",
        "credit.rejected",
        "credit.special_release_rejected",
    }
    for item in reversed(trace):
        if str(item.get("stage") or "") in terminal_stages:
            return _parse_datetime(item.get("ts"))
    return _parse_datetime(trace[-1].get("ts")) if trace else None


class AnalyticsService:
    def __init__(
        self,
        *,
        policy_path: str | None = None,
        case_root: str = "data/cases",
    ) -> None:
        self.policy = load_sla_policy(policy_path)
        self.repository = CaseRepository(case_root)

    @staticmethod
    def normalize_days(value: int | str | None) -> int:
        try:
            days = int(value or 30)
        except (TypeError, ValueError) as exc:
            raise ValueError("统计范围必须是天数。") from exc
        if days == 0:
            return 0
        if days not in {7, 30, 90}:
            raise ValueError("统计范围仅支持7天、30天、90天或全部。")
        return days

    def report(
        self, *, days: int | str | None = 30, now: datetime | None = None
    ) -> dict[str, Any]:
        current = _normalized_now(now)
        normalized_days = self.normalize_days(days)
        since = current - timedelta(days=normalized_days) if normalized_days else None
        cases = []
        for case in self.repository.list_cases():
            created_at = _parse_datetime(case.get("created_at"))
            if since and (not created_at or created_at < since):
                continue
            cases.append(case)

        intervals = [
            row
            for case in cases
            for row in task_intervals(case, now=current, policy=self.policy)
        ]
        completed_intervals = [row for row in intervals if row["completed"]]
        open_intervals = [row for row in intervals if not row["completed"]]
        completed_cases = []
        case_rows: list[dict[str, Any]] = []
        for case in cases:
            created_at = _parse_datetime(case.get("created_at"))
            terminal_at = _terminal_at(case)
            if terminal_at:
                completed_cases.append(case)
            duration = (
                max((terminal_at - created_at).total_seconds() / 3600, 0.0)
                if terminal_at and created_at
                else None
            )
            customer = dict(case.get("customer") or {})
            owner = dict(case.get("owner") or case.get("applicant") or {})
            sla = case_sla(case, now=current, policy=self.policy)
            case_rows.append(
                {
                    "case_id": str(case.get("case_id") or ""),
                    "customer_name": str(customer.get("customer_name") or ""),
                    "business_type": str(customer.get("business_type") or ""),
                    "owner": str(owner.get("display_name") or owner.get("username") or "未分配"),
                    "created_at": created_at.isoformat(timespec="seconds") if created_at else "",
                    "status": str(case.get("status") or ""),
                    "status_label": STATUS_LABELS.get(
                        str(case.get("status") or ""), str(case.get("status") or "")
                    ),
                    "completed": terminal_at is not None,
                    "successful": str(case.get("status") or "") in SUCCESS_STATUSES,
                    "completed_at": terminal_at.isoformat(timespec="seconds") if terminal_at else "",
                    "cycle_hours": round(duration, 1) if duration is not None else None,
                    "current_task": str(sla.get("waiting_for") or ""),
                    "current_task_label": str(sla.get("task_label") or ""),
                    "sla_state": str(sla.get("state") or "not_applicable"),
                    "sla_state_label": str(sla.get("state_label") or "无需处理"),
                }
            )

        node_stats = self._node_stats(intervals)
        backlog = self._backlog(open_intervals)
        trend = self._trend(cases, completed_intervals, current, since)
        completed_durations = [
            float(row["cycle_hours"])
            for row in case_rows
            if row["cycle_hours"] is not None
        ]
        on_time = sum(bool(row["met_sla"]) for row in completed_intervals)
        return {
            "range_days": normalized_days,
            "range_label": "全部" if not normalized_days else f"近{normalized_days}天",
            "generated_at": current.isoformat(timespec="seconds"),
            "policy_version": self.policy.get("version"),
            "metrics": {
                "case_total": len(cases),
                "completed_cases": len(completed_cases),
                "completion_rate": round(len(completed_cases) / len(cases), 4) if cases else None,
                "successful_cases": sum(row["successful"] for row in case_rows),
                "avg_cycle_hours": _average(completed_durations),
                "p95_cycle_hours": _percentile(completed_durations, 0.95),
                "completed_tasks": len(completed_intervals),
                "sla_on_time_rate": round(on_time / len(completed_intervals), 4)
                if completed_intervals
                else None,
                "current_backlog": len(open_intervals),
                "current_overdue": sum(not row["met_sla"] for row in open_intervals),
            },
            "nodes": node_stats,
            "backlog": backlog,
            "trend": trend,
            "cases": sorted(case_rows, key=lambda row: row["created_at"], reverse=True),
            "intervals": sorted(
                intervals, key=lambda row: (row["entered_at"], row["case_id"]), reverse=True
            ),
        }

    def _node_stats(self, intervals: list[dict[str, Any]]) -> list[dict[str, Any]]:
        rules = self.policy.get("tasks") or {}
        rows = []
        for task, rule_value in rules.items():
            rule = dict(rule_value or {})
            task_rows = [row for row in intervals if row["task"] == task]
            completed = [row for row in task_rows if row["completed"]]
            opened = [row for row in task_rows if not row["completed"]]
            durations = [float(row["duration_hours"]) for row in completed]
            on_time = sum(bool(row["met_sla"]) for row in completed)
            rows.append(
                {
                    "task": task,
                    "label": str(rule.get("label") or task),
                    "target_hours": float(rule.get("target_hours") or 0),
                    "completed": len(completed),
                    "on_time": on_time,
                    "attainment_rate": round(on_time / len(completed), 4) if completed else None,
                    "avg_hours": _average(durations),
                    "p95_hours": _percentile(durations, 0.95),
                    "open": len(opened),
                    "overdue": sum(not row["met_sla"] for row in opened),
                    "avg_open_hours": _average(
                        [float(row["duration_hours"]) for row in opened]
                    ),
                }
            )
        rows.sort(key=lambda row: (-row["overdue"], -row["open"], row["label"]))
        return rows

    @staticmethod
    def _backlog(intervals: list[dict[str, Any]]) -> list[dict[str, Any]]:
        groups: dict[str, dict[str, Any]] = {}
        for row in intervals:
            responsible = dict(row["responsible"])
            group = groups.setdefault(
                responsible["key"],
                {
                    "key": responsible["key"],
                    "label": responsible["label"],
                    "kind": responsible["kind"],
                    "total": 0,
                    "overdue": 0,
                    "tasks": defaultdict(int),
                },
            )
            group["total"] += 1
            group["overdue"] += int(not row["met_sla"])
            group["tasks"][row["task"]] += 1
        result = []
        for group in groups.values():
            group["tasks"] = dict(group["tasks"])
            result.append(group)
        result.sort(key=lambda row: (-row["overdue"], -row["total"], row["label"]))
        return result

    @staticmethod
    def _trend(
        cases: list[dict[str, Any]],
        completed_intervals: list[dict[str, Any]],
        current: datetime,
        since: datetime | None,
    ) -> list[dict[str, Any]]:
        dates = [
            parsed.date()
            for case in cases
            if (parsed := _parse_datetime(case.get("created_at")))
        ]
        start = since.date() if since else (min(dates) if dates else current.date())
        end = current.date()
        buckets: dict[str, dict[str, Any]] = {}
        cursor = start
        while cursor <= end:
            key = cursor.isoformat()
            buckets[key] = {
                "date": key,
                "created": 0,
                "completed": 0,
                "tasks_completed": 0,
                "tasks_on_time": 0,
            }
            cursor += timedelta(days=1)
        for case in cases:
            created_at = _parse_datetime(case.get("created_at"))
            if created_at and created_at.date().isoformat() in buckets:
                buckets[created_at.date().isoformat()]["created"] += 1
            terminal_at = _terminal_at(case)
            if terminal_at and terminal_at.date().isoformat() in buckets:
                buckets[terminal_at.date().isoformat()]["completed"] += 1
        for row in completed_intervals:
            exited_at = _parse_datetime(row.get("exited_at"))
            if exited_at and exited_at.date().isoformat() in buckets:
                bucket = buckets[exited_at.date().isoformat()]
                bucket["tasks_completed"] += 1
                bucket["tasks_on_time"] += int(bool(row["met_sla"]))
        return list(buckets.values())

    @staticmethod
    def export_csv(report: dict[str, Any]) -> bytes:
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(
            [
                "案件号",
                "客户",
                "业务类型",
                "节点",
                "责任队列",
                "进入时间",
                "离开时间",
                "处理时长（小时）",
                "目标时长（小时）",
                "状态",
                "SLA达标",
            ]
        )
        labels = {row["task"]: row["label"] for row in report["nodes"]}
        for row in report["intervals"]:
            writer.writerow(
                [
                    row["case_id"],
                    row["customer_name"],
                    row["business_type"],
                    labels.get(row["task"], row["task"]),
                    row["responsible"]["label"],
                    row["entered_at"],
                    row["exited_at"],
                    row["duration_hours"],
                    row["target_hours"],
                    "已完成" if row["completed"] else "进行中",
                    "是" if row["met_sla"] else "否",
                ]
            )
        return ("\ufeff" + output.getvalue()).encode("utf-8")

    @staticmethod
    def export_xlsx(report: dict[str, Any]) -> bytes:
        try:
            from openpyxl import Workbook
            from openpyxl.chart import BarChart, Reference
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        except ImportError as exc:  # pragma: no cover - production image includes documents extra
            raise RuntimeError("Excel导出组件未安装，请安装项目的documents依赖。") from exc

        workbook = Workbook()
        summary = workbook.active
        summary.title = "管理摘要"
        nodes = workbook.create_sheet("节点明细")
        cases = workbook.create_sheet("案件明细")
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True

        ink, primary, green, orange, red = (
            "222222",
            "FF385C",
            "167A52",
            "9A5B00",
            "C13515",
        )
        light_border = Border(bottom=Side(style="thin", color="E5E5E5"))

        def title(sheet: Any, text: str, end_column: int) -> None:
            sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=end_column)
            cell = sheet.cell(1, 1, text)
            cell.fill = PatternFill("solid", fgColor=primary)
            cell.font = Font(color="FFFFFF", bold=True, size=18)
            cell.alignment = Alignment(vertical="center")
            sheet.row_dimensions[1].height = 28
            sheet.row_dimensions[2].height = 8
            sheet.sheet_view.showGridLines = False

        def header(sheet: Any, row: int, columns: int) -> None:
            for cell in sheet[row][:columns]:
                cell.fill = PatternFill("solid", fgColor=ink)
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(vertical="center")
            sheet.row_dimensions[row].height = 24

        title(summary, f"审批时效管理报表 · {report['range_label']}", 8)
        summary["A4"] = "指标"
        summary["B4"] = "结果"
        summary["C4"] = "口径"
        header(summary, 4, 3)
        metrics = report["metrics"]
        metric_rows = [
            ("发起案件", metrics["case_total"], "按案件发起时间筛选"),
            ("已办结案件", metrics["completed_cases"], "含批准、拒绝和关闭"),
            ("案件办结率", metrics["completion_rate"], "已办结案件/发起案件"),
            ("平均端到端时长", metrics["avg_cycle_hours"], "已办结案件，小时"),
            ("P95端到端时长", metrics["p95_cycle_hours"], "已办结案件，小时"),
            ("已完成节点", metrics["completed_tasks"], "可重建进入和离开时间的节点"),
            ("节点SLA达标率", metrics["sla_on_time_rate"], "按已完成节点计算"),
            ("当前积压", metrics["current_backlog"], "当前仍在处理中的节点"),
            ("当前逾期", metrics["current_overdue"], "当前处理时长超过目标"),
        ]
        for row_index, row in enumerate(metric_rows, 5):
            for column_index, value in enumerate(row, 1):
                summary.cell(row_index, column_index, value)
            summary.cell(row_index, 1).font = Font(bold=True, color=ink)
            for cell in summary[row_index][:3]:
                cell.border = light_border
        summary["B7"].number_format = "0.0%"
        summary["B11"].number_format = "0.0%"
        summary["A15"] = "节点表现"
        summary["A15"].font = Font(bold=True, size=13, color=ink)
        summary.append([])
        node_headers = ["节点", "目标小时", "已完成", "达标率", "平均小时", "P95小时", "当前积压", "当前逾期"]
        for column, value in enumerate(node_headers, 1):
            summary.cell(16, column, value)
        header(summary, 16, len(node_headers))
        for row_index, row in enumerate(report["nodes"], 17):
            values = [
                row["label"], row["target_hours"], row["completed"],
                row["attainment_rate"], row["avg_hours"], row["p95_hours"],
                row["open"], row["overdue"],
            ]
            for column_index, value in enumerate(values, 1):
                summary.cell(row_index, column_index, value)
            summary.cell(row_index, 4).number_format = "0.0%"
        if report["nodes"]:
            chart = BarChart()
            chart.title = "当前节点积压"
            chart.y_axis.title = "待办数"
            chart.y_axis.numFmt = "0"
            chart.y_axis.majorUnit = 1
            chart.x_axis.title = "节点"
            data = Reference(summary, min_col=7, min_row=16, max_row=16 + len(report["nodes"]))
            categories = Reference(summary, min_col=1, min_row=17, max_row=16 + len(report["nodes"]))
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)
            chart.height = 7
            chart.width = 13
            chart.legend = None
            summary.add_chart(chart, "J4")
        summary.freeze_panes = "A5"
        summary.column_dimensions["A"].width = 22
        summary.column_dimensions["B"].width = 16
        summary.column_dimensions["C"].width = 32
        for column in "DEFGH":
            summary.column_dimensions[column].width = 14

        title(nodes, "节点处理明细", 11)
        node_detail_headers = [
            "案件号", "客户", "业务类型", "节点", "责任队列", "进入时间", "离开时间",
            "处理时长（小时）", "目标时长（小时）", "状态", "SLA达标",
        ]
        for column, value in enumerate(node_detail_headers, 1):
            nodes.cell(4, column, value)
        header(nodes, 4, len(node_detail_headers))
        labels = {row["task"]: row["label"] for row in report["nodes"]}
        for row_index, row in enumerate(report["intervals"], 5):
            values = [
                row["case_id"], row["customer_name"], row["business_type"],
                labels.get(row["task"], row["task"]), row["responsible"]["label"],
                _excel_datetime(row["entered_at"]), _excel_datetime(row["exited_at"]),
                row["duration_hours"], row["target_hours"],
                "已完成" if row["completed"] else "进行中", "是" if row["met_sla"] else "否",
            ]
            for column_index, value in enumerate(values, 1):
                nodes.cell(row_index, column_index, value)
            nodes.cell(row_index, 6).number_format = "yyyy-mm-dd hh:mm"
            nodes.cell(row_index, 7).number_format = "yyyy-mm-dd hh:mm"
            nodes.cell(row_index, 11).font = Font(color=green if row["met_sla"] else red, bold=True)
        nodes.freeze_panes = "A5"
        nodes.auto_filter.ref = f"A4:K{max(4, 4 + len(report['intervals']))}"
        node_widths = [18, 24, 12, 18, 22, 20, 20, 17, 17, 12, 12]
        for index, width in enumerate(node_widths, 1):
            nodes.column_dimensions[chr(64 + index)].width = width

        title(cases, "案件处理明细", 13)
        case_headers = [
            "案件号", "客户", "业务类型", "负责人", "发起时间", "状态", "已办结",
            "成功通过", "办结时间", "端到端时长（小时）", "当前节点", "当前SLA", "SLA状态码",
        ]
        for column, value in enumerate(case_headers, 1):
            cases.cell(4, column, value)
        header(cases, 4, len(case_headers))
        for row_index, row in enumerate(report["cases"], 5):
            values = [
                row["case_id"], row["customer_name"], row["business_type"], row["owner"],
                _excel_datetime(row["created_at"]), row["status_label"],
                "是" if row["completed"] else "否", "是" if row["successful"] else "否",
                _excel_datetime(row["completed_at"]), row["cycle_hours"],
                row["current_task_label"], row["sla_state_label"], row["sla_state"],
            ]
            for column_index, value in enumerate(values, 1):
                cases.cell(row_index, column_index, value)
            cases.cell(row_index, 5).number_format = "yyyy-mm-dd hh:mm"
            cases.cell(row_index, 9).number_format = "yyyy-mm-dd hh:mm"
            state_color = red if row["sla_state"] == "overdue" else orange if row["sla_state"] == "due_soon" else green
            cases.cell(row_index, 12).font = Font(color=state_color, bold=True)
        cases.freeze_panes = "A5"
        cases.auto_filter.ref = f"A4:M{max(4, 4 + len(report['cases']))}"
        case_widths = [18, 24, 12, 18, 20, 18, 12, 12, 20, 19, 18, 14, 15]
        for index, width in enumerate(case_widths, 1):
            cases.column_dimensions[chr(64 + index)].width = width

        for sheet in (summary, nodes, cases):
            for row in sheet.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(
                        vertical="center",
                        wrap_text=cell.column in {2, 3, 4, 5},
                    )
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
