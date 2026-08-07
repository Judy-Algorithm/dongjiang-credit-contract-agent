"""Versioned, isolated competition benchmark execution and exports."""

from __future__ import annotations

import csv
import io
import json
import statistics
import tempfile
import threading
import time
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from ..integrations import IntegrationBundle
from ..llm import ContractAIAssistant
from ..persistence import CaseRepository
from ..workflow import ActorContext, DongjiangWorkflowHarness
from ..workflow.dynamic import assert_plan_integrity


DEFAULT_SUITE_PATH = (
    Path(__file__).resolve().parents[1] / "config" / "benchmark_cases.json"
)


def _rate(numerator: int, denominator: int) -> float | None:
    return round(numerator / denominator, 4) if denominator else None


def _p95(values: list[float]) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    return round(ordered[max(0, int(len(ordered) * 0.95) - 1)], 2)


def _safe_value(value: Any) -> Any:
    if isinstance(value, float):
        return round(value, 4)
    if isinstance(value, list):
        return [_safe_value(item) for item in value]
    if isinstance(value, dict):
        return {str(key): _safe_value(item) for key, item in value.items()}
    return value


class BenchmarkService:
    """Run synthetic cases without touching production cases or integrations."""

    _run_lock = threading.Lock()

    def __init__(
        self,
        *,
        suite_path: str | Path = DEFAULT_SUITE_PATH,
        report_root: str | Path = "data/benchmarks",
    ) -> None:
        self.suite_path = Path(suite_path)
        self.report_root = Path(report_root)

    def load_suite(self) -> dict[str, Any]:
        payload = json.loads(self.suite_path.read_text(encoding="utf-8"))
        cases = payload.get("cases") or []
        if not payload.get("suite_id") or not payload.get("version"):
            raise ValueError("评测基准缺少suite_id或version。")
        if not isinstance(cases, list) or not cases:
            raise ValueError("评测基准必须包含非空cases数组。")
        keys = [str(item.get("case_key") or "") for item in cases]
        if any(not key for key in keys) or len(keys) != len(set(keys)):
            raise ValueError("评测案例case_key必须存在且唯一。")
        return payload

    def summary(self) -> dict[str, Any]:
        suite = self.load_suite()
        latest_path = self.report_root / "latest.json"
        latest: dict[str, Any] | None = None
        if latest_path.is_file():
            try:
                parsed = json.loads(latest_path.read_text(encoding="utf-8"))
                if isinstance(parsed, dict):
                    latest = parsed
            except (OSError, json.JSONDecodeError):
                latest = None
        return {
            "suite": self._suite_view(suite),
            "latest": latest,
            "status": "completed" if latest else "never_run",
        }

    def run(self, *, repeats: int = 2) -> dict[str, Any]:
        repeats = int(repeats)
        if repeats < 1 or repeats > 3:
            raise ValueError("重复次数必须在1至3次之间。")
        if not self._run_lock.acquire(blocking=False):
            raise RuntimeError("已有质量评测正在运行，请稍后刷新。")
        try:
            return self._run_locked(repeats=repeats)
        finally:
            self._run_lock.release()

    def _run_locked(self, *, repeats: int) -> dict[str, Any]:
        suite = self.load_suite()
        started_at = datetime.now(timezone.utc)
        rows_by_key: dict[str, list[dict[str, Any]]] = {
            str(item["case_key"]): [] for item in suite["cases"]
        }
        durations: list[float] = []
        with tempfile.TemporaryDirectory(prefix="dongjiang-benchmark-") as temp:
            root = Path(temp)
            for repeat_index in range(1, repeats + 1):
                for case in suite["cases"]:
                    result = self._run_case(
                        root / f"repeat-{repeat_index}" / str(case["case_key"]),
                        suite,
                        case,
                        repeat_index=repeat_index,
                    )
                    rows_by_key[str(case["case_key"])].append(result)
                    durations.append(float(result["duration_ms"]))
        cases = [
            self._merge_iterations(case, rows_by_key[str(case["case_key"])])
            for case in suite["cases"]
        ]
        report = self._report(
            suite,
            cases,
            repeats=repeats,
            started_at=started_at,
            durations=durations,
        )
        self._persist(report)
        return report

    def _run_case(
        self,
        root: Path,
        suite: dict[str, Any],
        case: dict[str, Any],
        *,
        repeat_index: int,
    ) -> dict[str, Any]:
        root.mkdir(parents=True, exist_ok=True)
        customer = deepcopy(suite.get("defaults") or {})
        customer.update(deepcopy(case.get("customer") or {}))
        started = time.perf_counter()
        actor = ActorContext("benchmark", ("system",), "benchmark", "质量评测")
        integrations = IntegrationBundle(audit_root=root / "integrations")
        checks: list[dict[str, Any]] = []
        error = ""
        state: dict[str, Any] = {}
        waiting_for: str | None = None
        try:
            with DongjiangWorkflowHarness(
                checkpoint_path=root / "workflow.sqlite",
                repository=CaseRepository(root / "cases"),
                vault_dir=root / "vault",
                inbox_dir=root / "inbox",
                output_dir=root / "output",
                evidence_dir=root / "evidence",
                archive_dir=root / "archive",
                execution_dir=root / "executions",
                integrations=integrations,
            ) as harness:
                harness.nodes.contract_ai = ContractAIAssistant(enabled=False)
                run = harness.start(
                    customer,
                    use_cached_credit=False,
                    actor=actor,
                )
                contract_text = self._contract_text(suite, case, customer)
                if contract_text and run.waiting_for == "credit_approval":
                    run = harness.resume(
                        run.case_id,
                        {"action": "approve", "comment": "合成基准自动批准信用建议"},
                        actor=actor,
                    )
                    if run.waiting_for == "special_release":
                        raise ValueError("合成基准合同案例意外触发信用锁定。")
                    run = harness.resume(
                        run.case_id,
                        {"action": "submit_contract", "contract_texts": [contract_text]},
                        actor=actor,
                    )
                    if run.waiting_for == "contract_approval":
                        run = harness.resume(
                            run.case_id,
                            {
                                "action": "approve",
                                "comment": "合成基准自动完成普通合同法务批准",
                            },
                            actor=actor,
                        )
                state = dict(run.state)
                waiting_for = run.waiting_for
        except Exception as exc:
            error = f"{type(exc).__name__}: {exc}"[:500]
        duration_ms = round((time.perf_counter() - started) * 1000, 2)
        if error:
            checks.append(
                self._check("execution", "案例执行完成", False, "无异常", error)
            )
        else:
            checks.extend(self._expected_checks(case, state, waiting_for))
            checks.extend(self._governance_checks(case, state))
        signature = self._signature(state, waiting_for) if state else {"error": error}
        return {
            "repeat": repeat_index,
            "duration_ms": duration_ms,
            "passed": bool(checks) and all(bool(item["passed"]) for item in checks),
            "checks": checks,
            "signature": signature,
            "error": error,
        }

    @staticmethod
    def _contract_text(
        suite: dict[str, Any], case: dict[str, Any], customer: dict[str, Any]
    ) -> str:
        contract = dict(case.get("contract") or {})
        if not contract:
            return ""
        template = str(contract.get("template") or "")
        source = str((suite.get("templates") or {}).get(template) or "")
        if not source:
            raise ValueError(f"评测合同模板不存在：{template}")
        values = {**dict(contract.get("values") or {})}
        values["customer_name"] = customer.get("customer_name") or "合成客户"
        return source.format(**values)

    @staticmethod
    def _check(
        category: str,
        label: str,
        passed: bool,
        expected: Any,
        actual: Any,
        *,
        metric: str = "",
    ) -> dict[str, Any]:
        return {
            "category": category,
            "label": label,
            "passed": bool(passed),
            "expected": _safe_value(expected),
            "actual": _safe_value(actual),
            "metric": metric,
        }

    def _expected_checks(
        self, case: dict[str, Any], state: dict[str, Any], waiting_for: str | None
    ) -> list[dict[str, Any]]:
        expected = dict(case.get("expected") or {})
        credit = dict(state.get("credit_assessment") or {})
        reviews = list(state.get("contract_reviews") or [])
        review = dict(reviews[0]) if reviews else {}
        findings = list(review.get("findings") or [])
        rule_ids = {str(item.get("rule_id") or "") for item in findings}
        checks: list[dict[str, Any]] = []
        credit_expected = dict(expected.get("credit") or {})
        for field, value in credit_expected.items():
            actual = (
                bool((credit.get("rating_resolution") or {}).get("requires_manual_review"))
                if field == "rating_requires_manual_review"
                else credit.get(field)
            )
            matched = (
                abs(float(actual or 0) - float(value)) < 0.01
                if isinstance(value, (int, float)) and not isinstance(value, bool)
                else actual == value
            )
            checks.append(
                self._check(
                    "credit",
                    self._credit_label(field),
                    matched,
                    value,
                    actual,
                    metric="credit",
                )
            )
        contract_expected = dict(expected.get("contract") or {})
        if "decision" in contract_expected:
            checks.append(
                self._check(
                    "contract",
                    "合同决策",
                    review.get("decision") == contract_expected["decision"],
                    contract_expected["decision"],
                    review.get("decision"),
                    metric="contract_decision",
                )
            )
        if "approval_route" in contract_expected:
            actual_route = state.get("approval_route") or review.get("approval_route")
            checks.append(
                self._check(
                    "route",
                    "审批路由",
                    actual_route == contract_expected["approval_route"],
                    contract_expected["approval_route"],
                    actual_route,
                    metric="route",
                )
            )
        for rule_id in contract_expected.get("required_rule_ids") or []:
            finding = next(
                (item for item in findings if item.get("rule_id") == rule_id), None
            )
            checks.append(
                self._check(
                    "rule",
                    f"命中规则 {rule_id}",
                    finding is not None,
                    "命中",
                    "命中" if finding else "未命中",
                    metric="rule_recall",
                )
            )
            if finding and (finding.get("clause_excerpt") or finding.get("evidence_query")):
                located = bool(finding.get("document_id") and finding.get("fragment_id"))
                checks.append(
                    self._check(
                        "evidence",
                        f"规则证据定位 {rule_id}",
                        located,
                        "文档及片段引用",
                        "已定位" if located else "未定位",
                        metric="evidence",
                    )
                )
        for rule_id in contract_expected.get("forbidden_rule_ids") or []:
            checks.append(
                self._check(
                    "rule",
                    f"不误报规则 {rule_id}",
                    rule_id not in rule_ids,
                    "不命中",
                    "命中" if rule_id in rule_ids else "未命中",
                    metric="false_positive_guard",
                )
            )
        workflow_expected = dict(expected.get("workflow") or {})
        if "waiting_for" in workflow_expected:
            checks.append(
                self._check(
                    "route",
                    "工作流等待节点",
                    waiting_for == workflow_expected["waiting_for"],
                    workflow_expected["waiting_for"],
                    waiting_for,
                    metric="route",
                )
            )
        return checks

    def _governance_checks(
        self, case: dict[str, Any], state: dict[str, Any]
    ) -> list[dict[str, Any]]:
        plans = list(state.get("workflow_plans") or [])
        audits = list(state.get("execution_audits") or [])
        verifications = [dict(state.get("credit_verification") or {})] + [
            dict(item) for item in state.get("contract_verifications") or []
        ]
        integrity_valid = True
        try:
            for plan in plans:
                assert_plan_integrity(plan)
        except Exception:
            integrity_valid = False
        audit_conformant = bool(audits) and all(
            item.get("status") == "conformant" for item in audits
        )
        verification_passed = bool(verifications) and all(
            item.get("status") == "passed" for item in verifications if item
        )
        logged = json.dumps(
            {
                "trace": state.get("trace") or [],
                "agent_runs": state.get("agent_runs") or [],
                "agent_task_results": state.get("agent_task_results") or [],
                "execution_audits": state.get("execution_audits") or [],
                "source_documents": state.get("source_documents") or [],
                "contract_facts": state.get("contract_facts") or [],
            },
            ensure_ascii=False,
        )
        sensitive_values = [str(item) for item in case.get("sensitive_values") or []]
        privacy_passed = all(value not in logged for value in sensitive_values)
        return [
            self._check(
                "agent",
                "冻结计划完整性",
                bool(plans) and integrity_valid,
                "全部有效",
                "全部有效" if plans and integrity_valid else "发现异常",
                metric="agent_conformance",
            ),
            self._check(
                "agent",
                "执行偏差审计",
                audit_conformant,
                "conformant",
                "conformant" if audit_conformant else "non_conformant",
                metric="agent_conformance",
            ),
            self._check(
                "agent",
                "独立核验",
                verification_passed,
                "passed",
                "passed" if verification_passed else "failed",
                metric="agent_conformance",
            ),
            self._check(
                "privacy",
                "敏感值未进入运行日志",
                privacy_passed,
                "无明文",
                "无明文" if privacy_passed else "发现明文",
                metric="privacy",
            ),
        ]

    @staticmethod
    def _signature(state: dict[str, Any], waiting_for: str | None) -> dict[str, Any]:
        credit = dict(state.get("credit_assessment") or {})
        reviews = list(state.get("contract_reviews") or [])
        return {
            "credit": {
                "score": credit.get("score"),
                "risk_level": credit.get("risk_level"),
                "approved_credit_limit": credit.get("approved_credit_limit"),
                "recommended_term_days": credit.get("recommended_term_days"),
                "requires_supplement": credit.get("requires_supplement"),
                "available_credit_amount": credit.get("available_credit_amount"),
            },
            "contract": [
                {
                    "decision": item.get("decision"),
                    "approval_route": item.get("approval_route"),
                    "rule_ids": sorted(
                        str(finding.get("rule_id") or "")
                        for finding in item.get("findings") or []
                    ),
                }
                for item in reviews
            ],
            "waiting_for": waiting_for,
            "agent": [
                {
                    "agent": plan.get("agent"),
                    "task_types": sorted(
                        str(task.get("task_type") or "")
                        for task in plan.get("tasks") or []
                    ),
                }
                for plan in state.get("workflow_plans") or []
            ],
        }

    def _merge_iterations(
        self, case: dict[str, Any], iterations: list[dict[str, Any]]
    ) -> dict[str, Any]:
        first = iterations[0]
        deterministic = all(
            item["signature"] == first["signature"] for item in iterations[1:]
        )
        checks = deepcopy(first["checks"])
        checks.append(
            self._check(
                "determinism",
                "重复执行结果一致",
                deterministic,
                "一致",
                "一致" if deterministic else "不一致",
                metric="determinism",
            )
        )
        return {
            "case_key": case["case_key"],
            "name": case["name"],
            "tags": list(case.get("tags") or []),
            "passed": all(bool(item["passed"]) for item in checks),
            "duration_ms": round(
                statistics.mean(float(item["duration_ms"]) for item in iterations), 2
            ),
            "checks": checks,
            "failed_checks": [
                item["label"] for item in checks if not bool(item["passed"])
            ],
            "actual": first["signature"],
            "repeat_count": len(iterations),
        }

    def _report(
        self,
        suite: dict[str, Any],
        cases: list[dict[str, Any]],
        *,
        repeats: int,
        started_at: datetime,
        durations: list[float],
    ) -> dict[str, Any]:
        checks = [check for case in cases for check in case["checks"]]

        def metric_rate(metric: str) -> float | None:
            rows = [item for item in checks if item.get("metric") == metric]
            return _rate(sum(bool(item["passed"]) for item in rows), len(rows))

        completed_at = datetime.now(timezone.utc)
        run_id = f"BRUN-{completed_at.strftime('%Y%m%d%H%M%S')}-{uuid4().hex[:6].upper()}"
        metrics = {
            "case_total": len(cases),
            "case_passed": sum(bool(item["passed"]) for item in cases),
            "case_pass_rate": _rate(
                sum(bool(item["passed"]) for item in cases), len(cases)
            ),
            "check_total": len(checks),
            "check_passed": sum(bool(item["passed"]) for item in checks),
            "regression_pass_rate": _rate(
                sum(bool(item["passed"]) for item in checks), len(checks)
            ),
            "credit_conclusion_rate": metric_rate("credit"),
            "contract_decision_rate": metric_rate("contract_decision"),
            "route_accuracy": metric_rate("route"),
            "rule_recall": metric_rate("rule_recall"),
            "false_positive_guard_rate": metric_rate("false_positive_guard"),
            "evidence_location_rate": metric_rate("evidence"),
            "agent_conformance_rate": metric_rate("agent_conformance"),
            "privacy_pass_rate": metric_rate("privacy"),
            "determinism_rate": metric_rate("determinism"),
            "average_latency_ms": round(statistics.mean(durations), 2),
            "p95_latency_ms": _p95(durations),
        }
        versions = self._versions(cases)
        return {
            "run_id": run_id,
            "status": "passed" if metrics["case_passed"] == len(cases) else "failed",
            "started_at": started_at.isoformat(timespec="seconds"),
            "completed_at": completed_at.isoformat(timespec="seconds"),
            "duration_ms": round(
                (completed_at - started_at).total_seconds() * 1000, 2
            ),
            "repeat_count": repeats,
            "suite": self._suite_view(suite),
            "runtime": {
                "isolation": "temporary_workspace",
                "external_ai": "disabled",
                "enterprise_integrations": "disabled",
                **versions,
            },
            "metrics": metrics,
            "cases": cases,
            "disclaimer": suite.get("disclaimer"),
        }

    @staticmethod
    def _versions(cases: list[dict[str, Any]]) -> dict[str, Any]:
        for case in cases:
            for plan in (case.get("actual") or {}).get("agent") or []:
                if plan.get("agent"):
                    return {"workflow_plan_version": "2.0"}
        return {"workflow_plan_version": "unknown"}

    def _persist(self, report: dict[str, Any]) -> None:
        self.report_root.mkdir(parents=True, exist_ok=True)
        encoded = json.dumps(report, ensure_ascii=False, indent=2)
        for target in (
            self.report_root / f"{report['run_id']}.json",
            self.report_root / "latest.json",
        ):
            temporary = target.with_suffix(target.suffix + ".tmp")
            temporary.write_text(encoded, encoding="utf-8")
            temporary.replace(target)

    @staticmethod
    def _suite_view(suite: dict[str, Any]) -> dict[str, Any]:
        return {
            "suite_id": suite.get("suite_id"),
            "version": suite.get("version"),
            "title": suite.get("title"),
            "policy_status": suite.get("policy_status"),
            "source": suite.get("source"),
            "case_count": len(suite.get("cases") or []),
            "tag_counts": {
                tag: sum(tag in (item.get("tags") or []) for item in suite.get("cases") or [])
                for tag in sorted(
                    {
                        str(tag)
                        for item in suite.get("cases") or []
                        for tag in item.get("tags") or []
                    }
                )
            },
        }

    @staticmethod
    def _credit_label(field: str) -> str:
        return {
            "risk_level": "信用风险等级",
            "approved_credit_limit": "建议授信额度",
            "recommended_term_days": "建议账期",
            "requires_supplement": "资料补充门禁",
            "available_credit_amount": "可用信用额度",
            "rating_requires_manual_review": "评级冲突人工复核",
        }.get(field, field)

    @staticmethod
    def export_csv(report: dict[str, Any]) -> bytes:
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(
            ["案例", "标签", "结果", "平均耗时ms", "检查类别", "检查项", "期望", "实际", "通过"]
        )
        for case in report.get("cases") or []:
            for check in case.get("checks") or []:
                writer.writerow(
                    [
                        case.get("name"),
                        "/".join(case.get("tags") or []),
                        "通过" if case.get("passed") else "失败",
                        case.get("duration_ms"),
                        check.get("category"),
                        check.get("label"),
                        json.dumps(check.get("expected"), ensure_ascii=False),
                        json.dumps(check.get("actual"), ensure_ascii=False),
                        "是" if check.get("passed") else "否",
                    ]
                )
        return ("\ufeff" + output.getvalue()).encode("utf-8")

    @staticmethod
    def export_xlsx(report: dict[str, Any]) -> bytes:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("Excel导出组件未安装，请安装documents依赖。") from exc
        workbook = Workbook()
        summary = workbook.active
        summary.title = "评测摘要"
        details = workbook.create_sheet("检查明细")
        cases_sheet = workbook.create_sheet("案例结果")
        summary.append(["东江信审与合同评审质量评测", report.get("run_id")])
        summary.append(["基准版本", (report.get("suite") or {}).get("version")])
        summary.append(["结果", report.get("status")])
        summary.append(["生成时间", report.get("completed_at")])
        summary.append([])
        summary.append(["指标", "结果"])
        for key, value in (report.get("metrics") or {}).items():
            summary.append([key, value])
        cases_sheet.append(["案例", "标签", "结果", "平均耗时ms", "失败检查"])
        details.append(["案例", "类别", "检查项", "期望", "实际", "通过"])
        for case in report.get("cases") or []:
            cases_sheet.append(
                [
                    case.get("name"),
                    "/".join(case.get("tags") or []),
                    "通过" if case.get("passed") else "失败",
                    case.get("duration_ms"),
                    "；".join(case.get("failed_checks") or []),
                ]
            )
            for check in case.get("checks") or []:
                details.append(
                    [
                        case.get("name"),
                        check.get("category"),
                        check.get("label"),
                        json.dumps(check.get("expected"), ensure_ascii=False),
                        json.dumps(check.get("actual"), ensure_ascii=False),
                        "是" if check.get("passed") else "否",
                    ]
                )
        for sheet in (summary, details, cases_sheet):
            sheet.freeze_panes = "A2"
            sheet.sheet_view.showGridLines = False
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="222222")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(vertical="center")
            for column in sheet.columns:
                letter = column[0].column_letter
                sheet.column_dimensions[letter].width = min(
                    48, max(12, max(len(str(cell.value or "")) for cell in column) + 2)
                )
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
