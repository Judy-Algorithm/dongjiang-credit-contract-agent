"""Case-level contract revision versions and Word redline generation."""

from __future__ import annotations

import copy
import hashlib
import json
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote
from uuid import uuid4

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree as ElementTree

from ..domain.models import utc_now
from ..ingestion import DocumentExtractor, location_label


WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
XML_NS = "http://www.w3.org/XML/1998/namespace"
W = f"{{{WORD_NS}}}"


REPLACEMENT_LIBRARY = {
    "DJ-CANCEL-WITHOUT-LIABILITY": (
        "订单取消：买方取消订单应至少提前30日书面通知，并补偿卖方已采购且无法合理转用的物料、"
        "在制品及已完成产品的合理成本；卖方应采取合理措施减少损失。"
    ),
    "DJ-INDIRECT-LOSS": (
        "违约责任：违约方仅对可证明的直接损失承担责任，不承担利润、商誉、市场份额等间接或后果性损失；"
        "累计赔偿责任不超过本合同已支付价款。"
    ),
    "LEGAL-UNLIMITED-LIABILITY": (
        "违约责任：除故意或重大过失外，违约方仅承担可证明的直接损失，累计赔偿责任不超过本合同已支付价款。"
    ),
    "DJ-AFFILIATE-SCOPE-SPILLOVER": (
        "责任主体：本合同项下权利义务仅由实际签约主体承担，不扩展至其关联公司、子公司、继受人或受让人，"
        "除非相关主体另行书面签署。"
    ),
    "DJ-JOINT-LIABILITY": "责任承担：各签约主体仅就自身行为和义务独立承担责任，任何一方均不承担连带责任。",
    "DJ-EXCLUSIVITY": "合作安排：双方合作为非排他性合作，不限制任何一方在本合同范围外与第三方开展业务。",
    "DJ-BANK-ACCEPTANCE": "付款方式：买方以银行转账方式支付到卖方指定账户。",
    "DJ-VMI-JIT-DELIVERY": (
        "交付与库存：VMI/JIT库存上限、风险转移、呆滞物料承担、交付计划及结算节点以双方确认的书面计划为准。"
    ),
    "DJ-WARRANTY-DUAL-LIMIT": (
        "质量保证：质保期以约定啤数或交付后约定期限两者先到为准，具体数值由双方在订单中书面确认。"
    ),
    "DJ-PRICE-FREEZE-CONTINUED-SUPPLY": (
        "价格调整：原材料、汇率或法规成本发生重大变化时双方应协商调价；协商期间持续供货义务最长不超过3个月。"
    ),
    "LEGAL-UNILATERAL-TERMINATION": (
        "解除与终止：一方发生重大违约且在收到书面催告后30日内未改正的，守约方可以书面通知解除合同；"
        "双方应结清已履行部分和合理已发生成本。"
    ),
    "LEGAL-IP-TRANSFER": (
        "知识产权：双方背景知识产权各自所有；项目成果的权属和许可范围由双方另行书面约定，不构成无偿整体转让。"
    ),
    "LEGAL-FOREIGN-JURISDICTION": "法律适用与争议解决：本合同适用中华人民共和国法律，争议由深圳市有管辖权的人民法院处理。",
    "LEGAL-HIGH-PENALTY": "违约金：违约金以受影响订单金额为计算基数，并以该订单金额的20%为累计上限。",
    "DJ-COMBINED-LIABILITY-OVER-50PCT": (
        "违约责任：违约金与损失赔偿金合计以受影响订单金额的50%为累计上限，且不包含任何间接或后果性损失。"
    ),
    "DJ-NO-AMOUNT-LIABILITY-OVER-HKD-1M": (
        "违约责任：本合同无具体交易金额时，违约金与损失赔偿金合计应低于港币100万元，且不承担间接或后果性损失。"
    ),
    "DJ-WARRANTY-DUAL-LIMIT-INCOMPLETE": (
        "质量保证：质保期以双方约定的啤数上限或交付后的期限上限两者先到为准。"
    ),
    "DJ-REPLACEMENT-WARRANTY-RESET": (
        "替代品质量保证：替代品或替换品继续适用原产品剩余质保期，替换不导致质保期重新起算或延长。"
    ),
    "DJ-SALES-COUNTRY-COMPLIANCE-SHIFT": (
        "法规合规：客户应书面告知产品销售目的国及适用的强制性要求；双方仅对各自控制范围内的法规识别和合规义务负责。"
    ),
    "DJ-IP-LICENSE-BOUNDARY-INCOMPLETE": (
        "知识产权：东江的背景知识产权仍归东江所有，仅就履行本合同之目的向客户提供有期限、免费且不可转让的许可。"
    ),
    "CREDIT-OVER-LIMIT": "付款条件：超出当前可用授信额度的部分应在发货前以预付款方式结清。",
    "TKP-SPECIAL-TERM": "付款及账期：货款账期为月结90天，起算日和对账要求按双方书面确认的结算规则执行。",
    "CREDIT-OVER-TERM": "付款及账期：货款账期不得超过正式信用审批确定的最长账期。",
}


def suggested_replacement(rule_id: str) -> str:
    return REPLACEMENT_LIBRARY.get(str(rule_id or ""), "")


def _safe_filename(value: str, fallback: str = "contract.docx") -> str:
    name = Path(str(value or "")).name
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip(" .")
    return name or fallback


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _artifact(path: Path, label: str) -> dict[str, Any]:
    return {
        "label": label,
        "name": path.name,
        "sha256": _sha256(path),
        "size_bytes": path.stat().st_size,
        "path": str(path.resolve()),
    }


def _style_generated_document(document: Document, *, author: str) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.top_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    document.core_properties.author = author or "东江集团"
    document.core_properties.last_modified_by = author or "东江集团"


def _text_to_docx(text: str, target: Path, *, author: str) -> None:
    document = Document()
    _style_generated_document(document, author=author)
    for index, line in enumerate(text.splitlines() or [text]):
        paragraph = document.add_paragraph()
        if index == 0 and line.strip() and len(line.strip()) <= 40:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(line)
            run.bold = True
            run.font.size = Pt(16)
            run.font.name = "Calibri"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        else:
            paragraph.add_run(line)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)


def _paragraph_text(paragraph: ElementTree.Element) -> str:
    return "".join(
        node.text or ""
        for node in paragraph.iter()
        if node.tag in {f"{W}t", f"{W}delText"}
    ).strip()


def _run_properties(paragraph: ElementTree.Element) -> ElementTree.Element | None:
    node = paragraph.find(f".//{W}rPr")
    return copy.deepcopy(node) if node is not None else None


def _append_run(
    parent: ElementTree.Element,
    text: str,
    properties: ElementTree.Element | None,
    *,
    deleted: bool,
    color: str,
) -> None:
    run = ElementTree.SubElement(parent, f"{W}r")
    rpr = copy.deepcopy(properties) if properties is not None else ElementTree.Element(f"{W}rPr")
    ElementTree.SubElement(rpr, f"{W}color", {f"{W}val": color})
    if deleted:
        ElementTree.SubElement(rpr, f"{W}strike")
    else:
        ElementTree.SubElement(rpr, f"{W}u", {f"{W}val": "single"})
    run.append(rpr)
    node = ElementTree.SubElement(run, f"{W}delText" if deleted else f"{W}t")
    if text[:1].isspace() or text[-1:].isspace():
        node.set(f"{{{XML_NS}}}space", "preserve")
    node.text = text


def _replace_paragraph_content(
    paragraph: ElementTree.Element,
    replacement: str,
    *,
    redline: bool,
    author: str,
    change_id: int,
    changed_at: str,
) -> None:
    original = _paragraph_text(paragraph)
    properties = _run_properties(paragraph)
    ppr = paragraph.find(f"{W}pPr")
    for child in list(paragraph):
        if child is not ppr:
            paragraph.remove(child)
    if not redline:
        run = ElementTree.SubElement(paragraph, f"{W}r")
        if properties is not None:
            run.append(properties)
        text = ElementTree.SubElement(run, f"{W}t")
        text.text = replacement
        return
    attributes = {
        f"{W}id": str(change_id),
        f"{W}author": author or "东江集团",
        f"{W}date": changed_at,
    }
    deleted = ElementTree.SubElement(paragraph, f"{W}del", attributes)
    _append_run(deleted, original, properties, deleted=True, color="C83D4F")
    inserted = ElementTree.SubElement(paragraph, f"{W}ins", {**attributes, f"{W}id": str(change_id + 1)})
    _append_run(inserted, replacement, properties, deleted=False, color="2463EB")


def _replace_table_cell_content(
    cell: ElementTree.Element,
    replacement: str,
    *,
    redline: bool,
    author: str,
    change_id: int,
    changed_at: str,
) -> None:
    paragraphs = cell.findall(f"./{W}p")
    if not paragraphs:
        paragraphs = [ElementTree.SubElement(cell, f"{W}p")]
    primary = paragraphs[0]
    original = "\n".join(
        value for value in (_paragraph_text(paragraph) for paragraph in paragraphs) if value
    )
    properties = _run_properties(primary)
    ppr = primary.find(f"{W}pPr")
    for child in list(primary):
        if child is not ppr:
            primary.remove(child)
    for paragraph in paragraphs[1:]:
        paragraph_properties = paragraph.find(f"{W}pPr")
        for child in list(paragraph):
            if child is not paragraph_properties:
                paragraph.remove(child)
    if not redline:
        run = ElementTree.SubElement(primary, f"{W}r")
        if properties is not None:
            run.append(properties)
        text = ElementTree.SubElement(run, f"{W}t")
        text.text = replacement
        return
    attributes = {
        f"{W}id": str(change_id),
        f"{W}author": author or "东江集团",
        f"{W}date": changed_at,
    }
    deleted = ElementTree.SubElement(primary, f"{W}del", attributes)
    _append_run(deleted, original, properties, deleted=True, color="C83D4F")
    inserted = ElementTree.SubElement(
        primary,
        f"{W}ins",
        {**attributes, f"{W}id": str(change_id + 1)},
    )
    _append_run(inserted, replacement, properties, deleted=False, color="2463EB")


def _patch_docx(
    source: Path,
    target: Path,
    replacements: dict[int, str],
    *,
    redline: bool,
    author: str,
    count_empty: bool = False,
    table_replacements: dict[tuple[int, int, int], str] | None = None,
) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    changed_at = datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")
    with zipfile.ZipFile(source, "r") as input_archive, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as output_archive:
        document_xml = ElementTree.fromstring(input_archive.read("word/document.xml"))
        current = 0
        change_id = 1
        for paragraph in document_xml.iter(f"{W}p"):
            if not count_empty and not _paragraph_text(paragraph):
                continue
            current += 1
            replacement = replacements.get(current)
            if replacement is None:
                continue
            _replace_paragraph_content(
                paragraph,
                replacement,
                redline=redline,
                author=author,
                change_id=change_id,
                changed_at=changed_at,
            )
            change_id += 2
        if current < max(replacements, default=0):
            raise ValueError("修订位置超出合同正文范围。")
        pending_table_replacements = dict(table_replacements or {})
        for table_number, table in enumerate(document_xml.iter(f"{W}tbl"), start=1):
            for row_number, row in enumerate(table.findall(f"./{W}tr"), start=1):
                column = 1
                for cell in row.findall(f"./{W}tc"):
                    properties = cell.find(f"./{W}tcPr")
                    span = properties.find(f"./{W}gridSpan") if properties is not None else None
                    try:
                        column_span = max(
                            1,
                            int(span.get(f"{W}val", "1") if span is not None else 1),
                        )
                    except (TypeError, ValueError):
                        column_span = 1
                    key = (table_number, row_number, column)
                    replacement = pending_table_replacements.pop(key, None)
                    if replacement is not None:
                        _replace_table_cell_content(
                            cell,
                            replacement,
                            redline=redline,
                            author=author,
                            change_id=change_id,
                            changed_at=changed_at,
                        )
                        change_id += 2
                    column += column_span
        if pending_table_replacements:
            raise ValueError("表格修订位置超出合同正文范围。")
        settings_xml: bytes | None = None
        if redline and "word/settings.xml" in input_archive.namelist():
            settings = ElementTree.fromstring(input_archive.read("word/settings.xml"))
            if settings.find(f"{W}trackRevisions") is None:
                settings.insert(0, ElementTree.Element(f"{W}trackRevisions"))
            settings_xml = ElementTree.tostring(settings, encoding="utf-8", xml_declaration=True)
        for item in input_archive.infolist():
            if item.filename == "word/document.xml":
                content = ElementTree.tostring(document_xml, encoding="utf-8", xml_declaration=True)
            elif item.filename == "word/settings.xml" and settings_xml is not None:
                content = settings_xml
            else:
                content = input_archive.read(item.filename)
            output_archive.writestr(item, content)


class ContractRevisionStore:
    def __init__(self, root: str | Path = "data/revisions") -> None:
        self.root = Path(root)

    def _case_root(self, case_id: str) -> Path:
        if not re.fullmatch(r"DJ-[A-Z0-9]+", case_id):
            raise ValueError("案件号格式无效。")
        return self.root / case_id

    def _manifest_path(self, case_id: str, revision_id: str) -> Path:
        if not re.fullmatch(r"REV-[A-Z0-9]+", revision_id):
            raise ValueError("修订版本号格式无效。")
        return self._case_root(case_id) / revision_id / "revision.json"

    @staticmethod
    def public(manifest: dict[str, Any]) -> dict[str, Any]:
        payload = {key: value for key, value in manifest.items() if key != "artifacts"}
        payload["decisions"] = [
            {
                key: value
                for key, value in decision.items()
                if key not in {"before", "after"}
            }
            for decision in manifest.get("decisions") or []
        ]
        payload["artifacts"] = {
            key: {name: value for name, value in artifact.items() if name != "path"}
            for key, artifact in (manifest.get("artifacts") or {}).items()
            if key in {"redline", "clean"}
        }
        return payload

    def list(self, case_id: str) -> list[dict[str, Any]]:
        case_root = self._case_root(case_id)
        rows: list[dict[str, Any]] = []
        if not case_root.exists():
            return rows
        for path in case_root.glob("REV-*/revision.json"):
            try:
                rows.append(self.public(json.loads(path.read_text(encoding="utf-8"))))
            except (OSError, ValueError, TypeError):
                continue
        return sorted(rows, key=lambda item: str(item.get("created_at") or ""), reverse=True)

    def get(self, case_id: str, revision_id: str) -> dict[str, Any]:
        path = self._manifest_path(case_id, revision_id)
        if not path.is_file():
            raise KeyError("合同修订版本不存在。")
        return json.loads(path.read_text(encoding="utf-8"))

    @staticmethod
    def _editable_findings(case: dict[str, Any], document_id: str) -> dict[str, dict[str, Any]]:
        rows: dict[str, dict[str, Any]] = {}
        for review in case.get("contract_reviews") or []:
            for finding in review.get("findings") or []:
                if finding.get("document_id") != document_id or not finding.get("fragment_id"):
                    continue
                key = f"{finding.get('rule_id')}:{document_id}:{finding.get('fragment_id')}"
                rows[key] = dict(finding)
        return rows

    def create(
        self,
        case: dict[str, Any],
        *,
        document_id: str,
        decisions: list[dict[str, Any]],
        actor: dict[str, Any],
    ) -> dict[str, Any]:
        case_id = str(case.get("case_id") or "")
        document = next(
            (
                dict(item)
                for item in case.get("source_documents") or []
                if item.get("document_id") == document_id and item.get("document_kind") == "contract"
            ),
            None,
        )
        if not document:
            raise KeyError("待修订合同不存在。")
        media_type = str(document.get("media_type") or "").lower()
        if media_type not in {"txt", "md", "docx"}:
            raise ValueError("当前仅支持修订 TXT、MD 和 DOCX 合同；PDF 请上传修改后的新版本。")
        source = Path(str(document.get("archived_path") or "")).resolve()
        expected_root = (Path("data/archive") / case_id).resolve()
        if expected_root not in source.parents or not source.is_file():
            raise PermissionError("合同归档路径无效。")

        editable = self._editable_findings(case, document_id)
        supplied = {str(item.get("finding_key") or ""): dict(item) for item in decisions}
        if not editable:
            raise ValueError("当前合同没有可定位的风险项，请上传人工修订版本。")
        if set(editable) != set(supplied):
            raise ValueError("必须逐项处理当前合同的全部可定位风险。")

        extracted = DocumentExtractor().extract(source)
        fragments = {item.fragment_id: item for item in extracted.fragments}
        replacements: dict[int, str] = {}
        table_replacements: dict[tuple[int, int, int], str] = {}
        resolved: list[dict[str, Any]] = []
        modified_fragments: set[str] = set()
        for key, finding in editable.items():
            decision = supplied[key]
            action = str(decision.get("action") or "")
            if action not in {"accept", "custom", "retain"}:
                raise ValueError("风险处置方式必须是采用建议、人工修改或保留说明。")
            fragment_id = str(finding.get("fragment_id") or "")
            fragment = fragments.get(fragment_id)
            if not fragment:
                raise ValueError("风险定位片段已经失效，请重新解析合同。")
            before = fragment.text
            reason = str(decision.get("reason") or "").strip()
            replacement = ""
            if action == "accept":
                replacement = suggested_replacement(str(finding.get("rule_id") or ""))
                if not replacement:
                    raise ValueError(f"{finding.get('title') or '该风险'}没有可直接采用的标准条款，请人工修改。")
            elif action == "custom":
                replacement = str(decision.get("replacement") or "").strip()
                if not replacement:
                    raise ValueError("人工修改必须填写替换后的完整条款。")
            elif not reason:
                raise ValueError("保留原条款必须填写原因。")
            if action != "retain":
                if fragment_id in modified_fragments:
                    raise ValueError("同一原文片段包含多个风险，请合并为一次人工修改，其余风险选择保留并说明。")
                modified_fragments.add(fragment_id)
                location = dict(fragment.location or {})
                if location.get("kind") == "word_table_cell":
                    table_position = (
                        int(location.get("table") or 0),
                        int(location.get("row") or 0),
                        int(location.get("column") or 0),
                    )
                    if min(table_position) <= 0:
                        raise ValueError("当前表格风险位置不支持自动修订。")
                    table_replacements[table_position] = replacement
                else:
                    position = int(location.get("paragraph") or location.get("line") or 0)
                    if position <= 0:
                        raise ValueError("当前风险位置不支持自动修订。")
                    replacements[position] = replacement
            resolved.append(
                {
                    "finding_key": key,
                    "rule_id": finding.get("rule_id"),
                    "title": finding.get("title"),
                    "action": action,
                    "reason": reason,
                    "fragment_id": fragment_id,
                    "location": dict(fragment.location or {}),
                    "location_label": location_label(fragment.location),
                    "before": before,
                    "after": replacement if action != "retain" else before,
                }
            )

        revision_id = f"REV-{uuid4().hex[:10].upper()}"
        revision_root = self._case_root(case_id) / revision_id
        revision_root.mkdir(parents=True, exist_ok=False)
        stem = Path(_safe_filename(str(document.get("name") or "contract"), "contract")).stem
        clean_path = revision_root / f"{stem}-{revision_id}-clean.docx"
        redline_path = revision_root / f"{stem}-{revision_id}-redline.docx"
        author = str(actor.get("display_name") or actor.get("actor_id") or "东江集团")

        if media_type == "docx":
            _patch_docx(
                source,
                clean_path,
                replacements,
                redline=False,
                author=author,
                table_replacements=table_replacements,
            )
            _patch_docx(
                source,
                redline_path,
                replacements,
                redline=True,
                author=author,
                table_replacements=table_replacements,
            )
        else:
            original_text = source.read_text(encoding="utf-8", errors="replace")
            revised_lines = original_text.splitlines()
            for number, replacement in replacements.items():
                if number > len(revised_lines):
                    raise ValueError("修订行号超出合同正文范围。")
                revised_lines[number - 1] = replacement
            clean_text = "\n".join(revised_lines)
            source_docx = revision_root / "source.docx"
            _text_to_docx(original_text, source_docx, author=author)
            _text_to_docx(clean_text, clean_path, author=author)
            _patch_docx(
                source_docx,
                redline_path,
                replacements,
                redline=True,
                author=author,
                count_empty=True,
            )
            source_docx.unlink(missing_ok=True)

        manifest = {
            "revision_id": revision_id,
            "case_id": case_id,
            "document_id": document_id,
            "source_name": document.get("name"),
            "source_sha256": document.get("sha256"),
            "status": "draft",
            "created_at": utc_now(),
            "created_by": {
                "actor_id": actor.get("actor_id"),
                "display_name": actor.get("display_name"),
            },
            "decisions": resolved,
            "artifacts": {
                "redline": _artifact(redline_path, "带修订痕迹版"),
                "clean": _artifact(clean_path, "清洁版"),
            },
            "submitted_at": "",
            "review_result": {},
        }
        self._manifest_path(case_id, revision_id).write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        return self.public(manifest)

    def artifact_path(self, case_id: str, revision_id: str, kind: str) -> tuple[Path, str]:
        if kind not in {"redline", "clean"}:
            raise KeyError("修订文件类型不存在。")
        manifest = self.get(case_id, revision_id)
        record = dict((manifest.get("artifacts") or {}).get(kind) or {})
        target = Path(str(record.get("path") or "")).resolve()
        expected = self._manifest_path(case_id, revision_id).parent.resolve()
        if expected not in target.parents or not target.is_file():
            raise PermissionError("修订文件路径无效。")
        return target, str(record.get("name") or target.name)

    def mark_submitted(self, case_id: str, revision_id: str, result: dict[str, Any]) -> dict[str, Any]:
        manifest = self.get(case_id, revision_id)
        manifest["status"] = "submitted"
        manifest["submitted_at"] = utc_now()
        manifest["review_result"] = result
        self._manifest_path(case_id, revision_id).write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        return self.public(manifest)


def content_disposition(filename: str) -> str:
    safe = _safe_filename(filename)
    return f"attachment; filename=contract.docx; filename*=UTF-8''{quote(safe)}"
