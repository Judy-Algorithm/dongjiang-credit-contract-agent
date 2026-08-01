"""Auditable, case-level contract translation drafts and bilingual exports."""

from __future__ import annotations

import hashlib
import json
import re
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any
from uuid import uuid4

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor
from lxml import etree as ElementTree

from ..domain.models import utc_now
from ..ingestion import DocumentExtractor, location_label
from ..llm import OpenAICompatibleGateway
from ..security import RedactionVault
from .revisions import W, _safe_filename, _sha256, _style_generated_document


LANGUAGES = {
    "zh": "中文",
    "en": "英文",
    "vi": "越南语",
    "ja": "日语",
    "es": "西班牙语",
}
_JSON_BLOCK = re.compile(r"```(?:json)?\s*(.*?)\s*```", re.IGNORECASE | re.DOTALL)
_TOKEN = re.compile(r"⟦[A-Z_]+[0-9a-f]{10}⟧")
_MAX_DOCUMENT_CHARS = 160_000
_MAX_BATCH_CHARS = 12_000
_MAX_BATCH_FRAGMENTS = 24


def _tokens(value: str) -> Counter[str]:
    return Counter(_TOKEN.findall(str(value or "")))


def _source_hash(value: str) -> str:
    return hashlib.sha256(str(value or "").encode("utf-8")).hexdigest()


def _target_languages(manifest: dict[str, Any]) -> list[str]:
    values = list(manifest.get("target_languages") or [])
    if not values and manifest.get("target_language"):
        values = [str(manifest["target_language"])]
    return [str(item) for item in values if str(item) in LANGUAGES]


def _entry_translations(
    entry: dict[str, Any], languages: list[str]
) -> dict[str, str]:
    rows = {
        str(language): str(text or "")
        for language, text in dict(entry.get("translations") or {}).items()
        if str(language) in LANGUAGES
    }
    if not rows and languages and entry.get("translated_text") is not None:
        rows[languages[0]] = str(entry.get("translated_text") or "")
    return rows


def _normalize_languages(
    target_languages: list[str] | tuple[str, ...] | None,
    target_language: str = "",
) -> list[str]:
    requested = list(target_languages or [])
    if not requested and target_language:
        requested = [target_language]
    languages: list[str] = []
    for item in requested:
        language = str(item or "").strip().lower()
        if language not in LANGUAGES:
            raise ValueError("目标语言必须是中文、英文、越南语、日语或西班牙语。")
        if language not in languages:
            languages.append(language)
    if not languages:
        raise ValueError("请至少选择一种目标语言。")
    return languages


def _artifact(path: Path) -> dict[str, Any]:
    return {
        "name": path.name,
        "sha256": _sha256(path),
        "size_bytes": path.stat().st_size,
        "path": str(path.resolve()),
    }


def _batch(entries: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    batches: list[list[dict[str, Any]]] = []
    current: list[dict[str, Any]] = []
    size = 0
    for entry in entries:
        entry_size = len(str(entry.get("source_text") or ""))
        if current and (
            len(current) >= _MAX_BATCH_FRAGMENTS
            or size + entry_size > _MAX_BATCH_CHARS
        ):
            batches.append(current)
            current = []
            size = 0
        current.append(entry)
        size += entry_size
    if current:
        batches.append(current)
    return batches


class ContractTranslator:
    prompt_version = "dongjiang-contract-translation-v1"

    def __init__(self, gateway: OpenAICompatibleGateway | None = None) -> None:
        self.gateway = gateway or OpenAICompatibleGateway()

    def translate(
        self,
        entries: list[dict[str, Any]],
        *,
        target_language: str,
    ) -> list[dict[str, Any]]:
        if target_language not in LANGUAGES:
            raise ValueError("目标语言必须是中文、英文、越南语、日语或西班牙语。")
        if not self.gateway.available:
            raise RuntimeError("文本模型尚未配置，无法生成合同译稿。")
        if not entries:
            raise ValueError("合同没有可翻译的文本片段。")
        translated: dict[str, str] = {}
        for batch in _batch(entries):
            request_entries = [
                {
                    "fragment_id": item["fragment_id"],
                    "text": item["source_text"],
                }
                for item in batch
            ]
            redacted_text = "⟦REDACTED_TEXT⟧\n" + json.dumps(
                {"segments": request_entries}, ensure_ascii=False
            )
            raw = self.gateway.analyze_redacted(
                redacted_text,
                self._instruction(target_language),
            )
            payload = self._parse(raw)
            rows = payload.get("translations")
            if not isinstance(rows, list):
                raise ValueError("模型翻译输出缺少 translations 数组。")
            expected = {item["fragment_id"]: item for item in batch}
            received: dict[str, str] = {}
            for row in rows:
                if not isinstance(row, dict):
                    continue
                fragment_id = str(row.get("fragment_id") or "")
                text = str(row.get("text") or "").strip()
                if fragment_id not in expected or fragment_id in received or not text:
                    continue
                if _tokens(text) != _tokens(str(expected[fragment_id]["source_text"])):
                    raise ValueError(f"模型改变了 {fragment_id} 的脱敏令牌，译稿已拒绝。")
                received[fragment_id] = text
            if set(received) != set(expected):
                missing = sorted(set(expected) - set(received))
                raise ValueError(f"模型未完整返回翻译片段：{', '.join(missing[:5])}")
            translated.update(received)
        return [
            {
                **entry,
                "translated_text": translated[str(entry["fragment_id"])],
            }
            for entry in entries
        ]

    @staticmethod
    def _instruction(target_language: str) -> str:
        return f"""将输入 JSON 中每个合同片段忠实翻译为{LANGUAGES[target_language]}。
只输出 JSON 对象，不要 Markdown：
{{"translations":[{{"fragment_id":"原片段ID","text":"翻译内容"}}]}}

要求：
1. 必须逐条返回，fragment_id、顺序和数量与输入完全一致，不得合并、拆分或遗漏。
2. ⟦...⟧ 是脱敏令牌，必须原样保留，不能翻译、删除或增加。
3. 保留数字、币种、日期、条款编号和法律含义，不补充原文没有的内容。
4. 使用正式合同语言；不输出分析、风险判断、批准意见或解释。"""

    @staticmethod
    def _parse(raw: str) -> dict[str, Any]:
        text = str(raw or "").strip()
        match = _JSON_BLOCK.search(text)
        if match:
            text = match.group(1).strip()
        payload = json.loads(text)
        if not isinstance(payload, dict):
            raise ValueError("模型翻译输出不是 JSON 对象。")
        return payload


class ContractTranslationStore:
    def __init__(
        self,
        root: str | Path = "data/translations",
        *,
        vault_dir: str | Path = "data/vault",
        translator: ContractTranslator | None = None,
    ) -> None:
        self.root = Path(root)
        self.vault_dir = Path(vault_dir)
        self.translator = translator or ContractTranslator()

    def _case_root(self, case_id: str) -> Path:
        if not re.fullmatch(r"DJ-[A-Z0-9]+", case_id):
            raise ValueError("案件号格式无效。")
        return self.root / case_id

    def _manifest_path(self, case_id: str, translation_id: str) -> Path:
        if not re.fullmatch(r"TR-[A-Z0-9]+", translation_id):
            raise ValueError("翻译版本号格式无效。")
        return self._case_root(case_id) / translation_id / "translation.json"

    @staticmethod
    def public(manifest: dict[str, Any]) -> dict[str, Any]:
        languages = _target_languages(manifest)
        labels = [LANGUAGES[item] for item in languages]
        return {
            "translation_id": manifest.get("translation_id"),
            "document_id": manifest.get("document_id"),
            "source_name": manifest.get("source_name"),
            "target_language": languages[0] if languages else "",
            "target_languages": languages,
            "target_language_label": " / ".join(labels),
            "target_language_labels": labels,
            "status": manifest.get("status"),
            "segment_count": len(manifest.get("entries") or []),
            "translation_count": len(manifest.get("entries") or []) * len(languages),
            "base_translation_id": manifest.get("base_translation_id"),
            "reused_translation_count": int(
                manifest.get("reused_translation_count") or 0
            ),
            "translated_translation_count": int(
                manifest.get("translated_translation_count") or 0
            ),
            "model": manifest.get("model"),
            "prompt_version": manifest.get("prompt_version"),
            "created_at": manifest.get("created_at"),
            "created_by": dict(manifest.get("created_by") or {}),
            "confirmed_at": manifest.get("confirmed_at"),
            "confirmed_by": dict(manifest.get("confirmed_by") or {}),
            "review_note": manifest.get("review_note"),
            "artifact": (
                {
                    key: value
                    for key, value in dict(manifest.get("artifact") or {}).items()
                    if key != "path"
                }
                if manifest.get("status") == "confirmed"
                else None
            ),
        }

    def list(self, case_id: str) -> list[dict[str, Any]]:
        case_root = self._case_root(case_id)
        if not case_root.exists():
            return []
        rows: list[dict[str, Any]] = []
        for path in case_root.glob("TR-*/translation.json"):
            try:
                rows.append(self.public(json.loads(path.read_text(encoding="utf-8"))))
            except (OSError, ValueError, TypeError):
                continue
        return sorted(rows, key=lambda item: str(item.get("created_at") or ""), reverse=True)

    def get(self, case_id: str, translation_id: str) -> dict[str, Any]:
        path = self._manifest_path(case_id, translation_id)
        if not path.is_file():
            raise KeyError("合同翻译版本不存在。")
        return json.loads(path.read_text(encoding="utf-8"))

    def detail(self, case_id: str, translation_id: str) -> dict[str, Any]:
        manifest = self.get(case_id, translation_id)
        vault = RedactionVault(case_id, self.vault_dir)
        languages = _target_languages(manifest)
        return {
            **self.public(manifest),
            "entries": [
                {
                    "fragment_id": item.get("fragment_id"),
                    "location": dict(item.get("location") or {}),
                    "location_label": item.get("location_label"),
                    "source_text": vault.restore(str(item.get("source_text") or "")),
                    "source_hash": item.get("source_hash") or _source_hash(
                        str(item.get("source_text") or "")
                    ),
                    "translations": {
                        language: vault.restore(text)
                        for language, text in _entry_translations(
                            item, languages
                        ).items()
                    },
                    "translation_sources": dict(
                        item.get("translation_sources") or {}
                    ),
                    "translated_text": vault.restore(str(
                        _entry_translations(item, languages).get(
                            languages[0] if languages else "", ""
                        )
                    )),
                }
                for item in manifest.get("entries") or []
            ],
        }

    def create(
        self,
        case: dict[str, Any],
        *,
        document_id: str,
        target_language: str = "",
        target_languages: list[str] | tuple[str, ...] | None = None,
        base_translation_id: str = "",
        actor: dict[str, Any],
    ) -> dict[str, Any]:
        case_id = str(case.get("case_id") or "")
        document = next(
            (
                dict(item)
                for item in case.get("source_documents") or []
                if item.get("document_id") == document_id
                and item.get("document_kind") == "contract"
                and item.get("parse_status") == "parsed"
            ),
            None,
        )
        if not document:
            raise KeyError("请选择已经成功解析的合同。")
        fragments = [
            {
                "fragment_id": str(item.get("fragment_id") or ""),
                "location": dict(item.get("location") or {}),
                "location_label": location_label(dict(item.get("location") or {})),
                "source_text": str(item.get("text") or "").strip(),
            }
            for item in document.get("fragments") or []
            if str(item.get("fragment_id") or "") and str(item.get("text") or "").strip()
        ]
        if not fragments:
            raise ValueError("合同没有可翻译的文本片段。")
        if sum(len(item["source_text"]) for item in fragments) > _MAX_DOCUMENT_CHARS:
            raise ValueError("合同正文超过16万字符，请拆分文件后翻译。")
        languages = _normalize_languages(target_languages, target_language)
        base = self._base_translation(
            case_id,
            languages,
            requested_id=base_translation_id,
        )
        base_languages = _target_languages(base) if base else []
        base_entries = {
            str(item.get("fragment_id") or ""): item
            for item in (base or {}).get("entries") or []
        }
        entries = []
        reused_count = 0
        translated_count = 0
        for fragment in fragments:
            prior = base_entries.get(str(fragment["fragment_id"])) or {}
            prior_source = str(prior.get("source_text") or "")
            prior_hash = str(prior.get("source_hash") or _source_hash(prior_source))
            same_source = prior_hash == _source_hash(str(fragment["source_text"]))
            prior_translations = _entry_translations(prior, base_languages)
            translations = {
                language: prior_translations[language]
                for language in languages
                if same_source and prior_translations.get(language)
            }
            sources = {language: "reused" for language in translations}
            reused_count += len(translations)
            entries.append({
                **fragment,
                "source_hash": _source_hash(str(fragment["source_text"])),
                "translations": translations,
                "translation_sources": sources,
            })
        for language in languages:
            pending = [
                entry for entry in entries if language not in entry["translations"]
            ]
            if not pending:
                continue
            translated = self.translator.translate(
                pending,
                target_language=language,
            )
            translated_by_id = {
                str(item["fragment_id"]): str(item["translated_text"])
                for item in translated
            }
            for entry in pending:
                fragment_id = str(entry["fragment_id"])
                entry["translations"][language] = translated_by_id[fragment_id]
                entry["translation_sources"][language] = "translated"
                translated_count += 1
        for entry in entries:
            entry["translated_text"] = entry["translations"][languages[0]]
        translation_id = f"TR-{uuid4().hex[:10].upper()}"
        root = self._case_root(case_id) / translation_id
        root.mkdir(parents=True, exist_ok=False)
        manifest = {
            "translation_id": translation_id,
            "case_id": case_id,
            "document_id": document_id,
            "source_name": document.get("name"),
            "source_media_type": document.get("media_type"),
            "source_sha256": document.get("sha256"),
            "source_path": document.get("archived_path"),
            "target_language": languages[0],
            "target_languages": languages,
            "base_translation_id": (base or {}).get("translation_id"),
            "reused_translation_count": reused_count,
            "translated_translation_count": translated_count,
            "status": "draft",
            "model": self.translator.gateway.model,
            "prompt_version": self.translator.prompt_version,
            "entries": entries,
            "created_at": utc_now(),
            "created_by": {
                "actor_id": actor.get("actor_id"),
                "display_name": actor.get("display_name"),
            },
            "confirmed_at": None,
            "confirmed_by": {},
            "review_note": "",
            "artifact": None,
        }
        self._write(manifest)
        return self.detail(case_id, translation_id)

    def _base_translation(
        self,
        case_id: str,
        languages: list[str],
        *,
        requested_id: str = "",
    ) -> dict[str, Any] | None:
        if requested_id:
            manifest = self.get(case_id, requested_id)
            if manifest.get("status") != "confirmed":
                raise ValueError("仅可复用已经人工确认的译稿。")
            return manifest
        case_root = self._case_root(case_id)
        candidates: list[dict[str, Any]] = []
        if case_root.exists():
            for path in case_root.glob("TR-*/translation.json"):
                try:
                    item = json.loads(path.read_text(encoding="utf-8"))
                except (OSError, ValueError, TypeError):
                    continue
                if item.get("status") != "confirmed":
                    continue
                if not set(languages).intersection(_target_languages(item)):
                    continue
                candidates.append(item)
        return max(
            candidates,
            key=lambda item: str(item.get("confirmed_at") or item.get("created_at") or ""),
            default=None,
        )

    def confirm(
        self,
        case_id: str,
        translation_id: str,
        *,
        entries: list[dict[str, Any]],
        review_note: str,
        actor: dict[str, Any],
    ) -> dict[str, Any]:
        manifest = self.get(case_id, translation_id)
        if manifest.get("status") != "draft":
            raise ValueError("该翻译版本已经确认，不能重复修改。")
        note = str(review_note or "").strip()
        if not note:
            raise ValueError("确认译稿前必须填写人工复核说明。")
        languages = _target_languages(manifest)
        supplied: dict[str, dict[str, str]] = {}
        for row in entries:
            if not isinstance(row, dict):
                continue
            fragment_id = str(row.get("fragment_id") or "")
            translations = {
                str(language): str(text or "").strip()
                for language, text in dict(row.get("translations") or {}).items()
                if str(language) in languages
            }
            if not translations and languages:
                translations[languages[0]] = str(
                    row.get("translated_text") or ""
                ).strip()
            supplied[fragment_id] = translations
        expected = {
            str(item.get("fragment_id") or ""): item
            for item in manifest.get("entries") or []
        }
        if set(supplied) != set(expected):
            raise ValueError("必须逐段确认当前译稿的全部内容。")
        vault = RedactionVault(case_id, self.vault_dir)
        for fragment_id, item in expected.items():
            if set(supplied[fragment_id]) != set(languages):
                raise ValueError(f"{fragment_id} 必须确认全部目标语言译文。")
            confirmed_translations: dict[str, str] = {}
            for language in languages:
                translated = supplied[fragment_id][language]
                if not translated or len(translated) > 20_000:
                    raise ValueError(
                        f"{fragment_id} 的{LANGUAGES[language]}译文为空或过长。"
                    )
                redacted = vault.redact(translated)
                if _tokens(redacted) != _tokens(str(item.get("source_text") or "")):
                    raise ValueError(
                        f"{fragment_id} 的{LANGUAGES[language]}译文未完整保留合同敏感字段。"
                    )
                confirmed_translations[language] = redacted
            item["translations"] = confirmed_translations
            item["translated_text"] = confirmed_translations[languages[0]]
        vault.persist_local()
        artifact_path = self._export_multilingual(manifest, vault)
        manifest.update(
            {
                "status": "confirmed",
                "entries": list(expected.values()),
                "confirmed_at": utc_now(),
                "confirmed_by": {
                    "actor_id": actor.get("actor_id"),
                    "display_name": actor.get("display_name"),
                },
                "review_note": note[:1000],
                "artifact": _artifact(artifact_path),
            }
        )
        self._write(manifest)
        return self.detail(case_id, translation_id)

    def artifact_path(self, case_id: str, translation_id: str) -> tuple[Path, str]:
        manifest = self.get(case_id, translation_id)
        if manifest.get("status") != "confirmed":
            raise ValueError("译稿尚未人工确认，不能下载正式双语稿。")
        artifact = dict(manifest.get("artifact") or {})
        target = Path(str(artifact.get("path") or "")).resolve()
        expected_root = (self._case_root(case_id) / translation_id).resolve()
        if expected_root not in target.parents or not target.is_file():
            raise PermissionError("双语合同归档路径无效。")
        return target, str(artifact.get("name") or target.name)

    def _write(self, manifest: dict[str, Any]) -> None:
        target = self._manifest_path(
            str(manifest.get("case_id") or ""),
            str(manifest.get("translation_id") or ""),
        )
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = target.with_suffix(".json.tmp")
        temporary.write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        temporary.replace(target)

    def _export_bilingual(
        self,
        manifest: dict[str, Any],
        vault: RedactionVault,
    ) -> Path:
        return self._export_multilingual(manifest, vault)

    def _export_multilingual(
        self,
        manifest: dict[str, Any],
        vault: RedactionVault,
    ) -> Path:
        root = self._case_root(str(manifest["case_id"])) / str(
            manifest["translation_id"]
        )
        source_name = _safe_filename(str(manifest.get("source_name") or "contract"))
        languages = _target_languages(manifest)
        language_suffix = "-".join(languages)
        target = root / f"{Path(source_name).stem}-{language_suffix}-multilingual.docx"
        translations = {
            str(item["fragment_id"]): {
                language: vault.restore(text)
                for language, text in _entry_translations(item, languages).items()
            }
            for item in manifest.get("entries") or []
        }
        source = Path(str(manifest.get("source_path") or "")).resolve()
        expected_archive = (Path("data/archive") / str(manifest["case_id"])).resolve()
        if expected_archive not in source.parents or not source.is_file():
            raise PermissionError("合同原件归档路径无效。")
        if source.suffix.lower() == ".docx":
            self._export_docx(source, target, translations, languages)
        else:
            extracted = DocumentExtractor().extract(source)
            source_by_id = {item.fragment_id: item.text for item in extracted.fragments}
            self._export_generated(
                target,
                source_name=source_name,
                entries=list(manifest.get("entries") or []),
                source_by_id=source_by_id,
                translations=translations,
                target_languages=languages,
            )
        return target

    @staticmethod
    def _translation_paragraph(text: str, language: str) -> ElementTree.Element:
        paragraph = ElementTree.Element(f"{W}p")
        ppr = ElementTree.SubElement(paragraph, f"{W}pPr")
        ElementTree.SubElement(ppr, f"{W}spacing", {f"{W}before": "40", f"{W}after": "120"})
        run = ElementTree.SubElement(paragraph, f"{W}r")
        rpr = ElementTree.SubElement(run, f"{W}rPr")
        ElementTree.SubElement(rpr, f"{W}color", {f"{W}val": "4B5563"})
        ElementTree.SubElement(rpr, f"{W}i")
        node = ElementTree.SubElement(run, f"{W}t")
        node.text = f"[{LANGUAGES[language]}译文] {text}"
        return paragraph

    @staticmethod
    def _translation_notice() -> ElementTree.Element:
        paragraph = ElementTree.Element(f"{W}p")
        ppr = ElementTree.SubElement(paragraph, f"{W}pPr")
        ElementTree.SubElement(
            ppr,
            f"{W}shd",
            {f"{W}fill": "FFF4E5", f"{W}val": "clear"},
        )
        run = ElementTree.SubElement(paragraph, f"{W}r")
        rpr = ElementTree.SubElement(run, f"{W}rPr")
        ElementTree.SubElement(rpr, f"{W}b")
        ElementTree.SubElement(rpr, f"{W}color", {f"{W}val": "9A4D00"})
        node = ElementTree.SubElement(run, f"{W}t")
        node.text = "翻译提示：本文件包含机器翻译并经人工确认，法律效力仍以签署原文为准。"
        return paragraph

    @classmethod
    def _export_docx(
        cls,
        source: Path,
        target: Path,
        translations: dict[str, dict[str, str]],
        target_languages: list[str],
    ) -> None:
        with zipfile.ZipFile(source, "r") as input_archive, zipfile.ZipFile(
            target, "w", zipfile.ZIP_DEFLATED
        ) as output_archive:
            document_xml = ElementTree.fromstring(input_archive.read("word/document.xml"))
            table_paragraphs = {
                paragraph
                for table in document_xml.iter(f"{W}tbl")
                for paragraph in table.iter(f"{W}p")
            }
            paragraphs = list(document_xml.iter(f"{W}p"))
            paragraph_number = 0
            for paragraph in paragraphs:
                text = "".join(
                    node.text or "" for node in paragraph.iter(f"{W}t")
                ).strip()
                if not text:
                    continue
                paragraph_number += 1
                fragment_id = f"paragraph-{paragraph_number}"
                translated = translations.get(fragment_id) or {}
                if translated and paragraph not in table_paragraphs:
                    anchor = paragraph
                    for language in target_languages:
                        if translated.get(language):
                            translation = cls._translation_paragraph(
                                translated[language], language
                            )
                            anchor.addnext(translation)
                            anchor = translation
            for table_number, table in enumerate(document_xml.iter(f"{W}tbl"), start=1):
                for row_number, row in enumerate(table.findall(f"./{W}tr"), start=1):
                    column = 1
                    for cell in row.findall(f"./{W}tc"):
                        properties = cell.find(f"./{W}tcPr")
                        span = properties.find(f"./{W}gridSpan") if properties is not None else None
                        try:
                            column_span = max(
                                1,
                                int(span.get(f"{W}val", "1") if span is not None else 1),
                            )
                        except (TypeError, ValueError):
                            column_span = 1
                        fragment_id = (
                            f"table-{table_number}-row-{row_number}-column-{column}"
                        )
                        translated = translations.get(fragment_id) or {}
                        for language in target_languages:
                            if translated.get(language):
                                cell.append(cls._translation_paragraph(
                                    translated[language], language
                                ))
                        column += column_span
            body = document_xml.find(f"./{W}body")
            if body is not None:
                first_content = next(
                    (item for item in body if item.tag != f"{W}sectPr"),
                    None,
                )
                if first_content is None:
                    body.insert(0, cls._translation_notice())
                else:
                    first_content.addprevious(cls._translation_notice())
            for item in input_archive.infolist():
                content = (
                    ElementTree.tostring(
                        document_xml, encoding="utf-8", xml_declaration=True
                    )
                    if item.filename == "word/document.xml"
                    else input_archive.read(item.filename)
                )
                output_archive.writestr(item, content)

    @staticmethod
    def _export_generated(
        target: Path,
        *,
        source_name: str,
        entries: list[dict[str, Any]],
        source_by_id: dict[str, str],
        translations: dict[str, dict[str, str]],
        target_languages: list[str],
    ) -> None:
        document = Document()
        _style_generated_document(document, author="东江集团")
        title = document.add_heading("合同多语言对照稿", level=0)
        title.alignment = 1
        document.add_paragraph(f"原文件：{source_name}")
        warning = document.add_paragraph(
            "本文件包含机器翻译内容，已经人工确认，但法律效力仍以签署原文为准。"
        )
        warning.runs[0].bold = True
        warning.runs[0].font.color.rgb = RGBColor(180, 60, 50)
        for index, entry in enumerate(entries, start=1):
            if index > 1:
                document.add_paragraph().add_run().add_break(WD_BREAK.LINE)
            label = document.add_paragraph(str(entry.get("location_label") or "文档原文"))
            label.runs[0].bold = True
            label.runs[0].font.size = Pt(9)
            source_paragraph = document.add_paragraph(
                source_by_id.get(str(entry.get("fragment_id") or ""), "")
            )
            source_paragraph.paragraph_format.space_after = Pt(3)
            values = translations.get(str(entry["fragment_id"])) or {}
            for language in target_languages:
                if not values.get(language):
                    continue
                translated = document.add_paragraph(
                    f"[{LANGUAGES[language]}译文] {values[language]}"
                )
                translated.runs[0].italic = True
                translated.runs[0].font.color.rgb = RGBColor(75, 85, 99)
        target.parent.mkdir(parents=True, exist_ok=True)
        document.save(target)
